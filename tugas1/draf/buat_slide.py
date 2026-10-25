"""Membuat slide presentasi Tugas 1 (Analisis Distribusi Effort — Phoenix Pay), 5 menit.

7 slide dipresentasikan (±45–60 detik per slide) + 2 slide cadangan (daftar istilah, daftar pustaka).
Versi untuk audiens awam: istilah diganti kata sehari-hari (istilah asli dalam kurung).
Logo Amikom di slide judul (kanan atas) dan pojok kiri bawah slide lainnya.

Jalankan dengan venv mata kuliah (berisi python-pptx):
    ../../.venv/bin/python buat_slide.py

Hasil: ../hasil/Tugas1_DistribusiEffort_Phoenix_Slide.pptx
Butuh: ../hasil/*_slide.png (buat dulu dengan buat_diagram.py) dan ../sumber/logo_amikom.png.
Semua angka uang dalam dolar Kanada (C$).
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TUGAS = Path(__file__).resolve().parent.parent
HASIL = TUGAS / "hasil"
LOGO = TUGAS / "sumber" / "logo_amikom.png"
OUT = HASIL / "Tugas1_DistribusiEffort_Phoenix_Slide.pptx"

PUTIH = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TINTA = RGBColor(0x1F, 0x1F, 0x1F)
ABU = RGBColor(0x52, 0x51, 0x4E)
REDUP = RGBColor(0x89, 0x87, 0x81)
BIRU = RGBColor(0x2A, 0x78, 0xD6)
BIRU_MUDA = RGBColor(0xEA, 0xF2, 0xFC)
MERAH = RGBColor(0xD0, 0x3B, 0x3B)
MERAH_MUDA = RGBColor(0xFB, 0xEA, 0xEA)
ABU_MUDA = RGBColor(0xF3, 0xF3, 0xF1)
JUDUL_FONT, ISI_FONT = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
KOSONG = prs.slide_layouts[6]
nomor = {"n": 0}


def teks(slide, x, y, w, h, isi, size=16, color=TINTA, bold=False, italic=False,
         font=ISI_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    """isi: str, atau list paragraf; tiap paragraf str atau list run (teks, {opsi})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.05))
    paragraf = isi if isinstance(isi, list) else [isi]
    for i, par in enumerate(paragraf):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = par if isinstance(par, list) else [(par, {})]
        for t, opt in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", italic)
            f.color.rgb = opt.get("color", color)
    return tb


def kotak(slide, x, y, w, h, fill=BIRU_MUDA, bentuk=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08):
    s = slide.shapes.add_shape(bentuk, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if bentuk == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def lingkaran_nomor(slide, x, y, angka, d=0.55, fill=BIRU, size=18):
    c = kotak(slide, x, y, d, d, fill=fill, bentuk=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(angka)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = ISI_FONT, Pt(size), True, PUTIH
    return c


def istilah(slide, x, y, w, h, pasangan, size=13):
    """Kotak 'Istilah di slide ini': pasangan = [(istilah, penjelasan), ...]."""
    kotak(slide, x, y, w, h, fill=ABU_MUDA, radius=0.06)
    baris = [[("Istilah di slide ini", {"bold": True, "color": NAVY})]]
    for t, d in pasangan:
        baris.append([(t + ": ", {"bold": True}), (d, {})])
    teks(slide, x + 0.15, y + 0.08, w - 0.3, h - 0.16, baris, size=size, color=TINTA, space_after=3)


def slide_baru(judul=None, catatan="", judul_slide=False):
    s = prs.slides.add_slide(KOSONG)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PUTIH
    nomor["n"] += 1
    if judul:
        teks(s, 0.6, 0.35, 12.1, 0.9, judul, size=32, bold=True, color=NAVY, font=JUDUL_FONT,
             anchor=MSO_ANCHOR.MIDDLE)
    if judul_slide:
        s.shapes.add_picture(str(LOGO), Inches(11.55), Inches(0.45), height=Inches(1.1))
    else:
        s.shapes.add_picture(str(LOGO), Inches(0.35), Inches(6.98), height=Inches(0.4))
        teks(s, 12.2, 7.0, 0.8, 0.35, str(nomor["n"]), size=11, color=REDUP, align=PP_ALIGN.RIGHT)
    if catatan:
        s.notes_slide.notes_text_frame.text = catatan
    return s


# =============== 1. Judul ===============
s = slide_baru(judul_slide=True, catatan=(
    "Selamat pagi/siang. Saya membahas mengapa proyek sistem gaji pemerintah Kanada, Phoenix, gagal. "
    "Fokusnya: bagaimana tenaga dan waktu proyek, atau effort, dibagi ke lima kelompok kegiatan pengelolaan "
    "proyek."))
teks(s, 0.9, 1.3, 10.4, 0.5, "TUGAS 1 — IT PROJECT MANAGEMENT", size=16, bold=True, color=BIRU)
teks(s, 0.9, 1.85, 10.4, 1.5, ["Analisis Distribusi Effort Proyek:", "Studi Kasus Phoenix Pay System (Kanada)"],
     size=32, bold=True, color=NAVY, font=JUDUL_FONT, space_after=0)
teks(s, 0.9, 3.45, 11.5, 0.8,
     "Mengapa sistem gaji untuk 290.000 pegawai pemerintah gagal, dilihat dari pembagian tenaga dan waktu (effort) proyek",
     size=18, italic=True, color=ABU)
for i, (nama, arti) in enumerate([("Initiating", "memulai"), ("Planning", "merencanakan"),
                                  ("Executing", "mengerjakan"), ("Monitoring &\nControlling", "memeriksa"),
                                  ("Closing", "menutup")]):
    x = 0.9 + i * 1.75
    lingkaran_nomor(s, x, 4.75, i + 1, d=0.5, size=16)
    teks(s, x - 0.4, 5.3, 1.3, 0.9, [nama, [(arti, {"italic": True, "color": BIRU})]],
         size=12, color=ABU, align=PP_ALIGN.CENTER, space_after=0)
teks(s, 9.9, 4.7, 3.0, 1.8, [
    [("Nama: ", {"bold": True}), ("EKA PUTRA KURNIA", {})],
    [("NIM: ", {"bold": True}), ("25.55.2927", {})],
    "S2 Informatika — Universitas Amikom Yogyakarta",
    "Dosen: Melwin Syafrizal, Ph.D.",
], size=13, color=TINTA, space_after=2)

# =============== 2. Kasus dan alasan ===============
s = slide_baru("Phoenix Pay: kasusnya dan mengapa dipilih", catatan=(
    "Phoenix adalah sistem gaji untuk 290 ribu pegawai di 101 lembaga Kanada, dibangun IBM. Uang yang disetujui "
    "155 juta dolar Kanada, padahal biayanya ternyata 274 juta. Sistem tetap dipakai 2016; setahun kemudian lebih dari "
    "520 juta dolar gaji salah atau terlambat. Saya memilih kasus ini karena diperiksa resmi oleh lembaga pemeriksa "
    "negara Kanada, setara BPK, yang memakai buku PMBOK seperti materi kuliah kita."))
angka = [
    ("290.000", "pegawai di 101 lembaga digaji lewat Phoenix", BIRU, BIRU_MUDA),
    ("C$155 → 274 jt", "uang disetujui vs biaya yang ternyata dibutuhkan", BIRU, BIRU_MUDA),
    ("> C$520 jt", "gaji salah atau terlambat (Juni 2017)", MERAH, MERAH_MUDA),
    ("51%", "pegawai menerima slip gaji salah (April 2017)", MERAH, MERAH_MUDA),
]
for i, (big, small, warna, latar) in enumerate(angka):
    x = 0.6 + (i % 2) * 3.95
    y = 1.45 + (i // 2) * 2.5
    kotak(s, x, y, 3.75, 2.3, fill=latar)
    teks(s, x + 0.25, y + 0.25, 3.3, 0.9, big, size=30, bold=True, color=warna, font=JUDUL_FONT)
    teks(s, x + 0.25, y + 1.15, 3.3, 1.0, small, size=16, color=TINTA)
kotak(s, 8.65, 1.45, 4.1, 4.8, fill=PUTIH, line=BIRU)
teks(s, 8.9, 1.65, 3.6, 0.5, "Mengapa kasus ini?", size=19, bold=True, color=BIRU)
teks(s, 8.9, 2.25, 3.6, 3.9, [
    [("Diperiksa resmi ", {"bold": True}), ("oleh OAG, lembaga pemeriksa negara Kanada (seperti BPK), 2018.", {})],
    [("Kerangka sama dengan kuliah: ", {"bold": True}), ("pemeriksa memakai PMBOK Guide.", {})],
    [("Angkanya konkret: ", {"bold": True}), ("fitur dibuang, pengujian gagal, jadwal dipadatkan.", {})],
], size=16, color=TINTA, space_after=12)
teks(s, 0.6, 6.4, 8.0, 0.45, "C$ = dolar Kanada. Pembangun sistem: IBM. Sumber: OAG Kanada (2018).",
     size=13, color=REDUP)

# =============== 3. Lini waktu ===============
s = slide_baru("Lini waktu: pemeriksaan yang dilewati", catatan=(
    "Titik biru adalah kejadian proyek; tanda silang merah adalah pemeriksaan yang dilewati. 2012: biaya membengkak "
    "tetapi rencana tidak disusun ulang, lebih dari 100 fitur dibuang. 2015: uji coba dibatalkan. Panel kanan "
    "memperbesar awal 2016: saran untuk meluncurkan bertahap diabaikan, sistem dilaporkan siap tanpa keputusan tertulis, "
    "dan jeda antartahap hanya dua bulan dari rencana tujuh."))
s.shapes.add_picture(str(HASIL / "Tugas1_DistribusiEffort_Phoenix_Diagram2_LiniWaktu_slide.png"),
                     Inches(0.97), Inches(1.3), width=Inches(11.4))

# =============== 4. Pembagian effort ===============
s = slide_baru("Pembagian effort: Phoenix vs angka pembanding", catatan=(
    "Pengelolaan proyek punya lima kelompok kegiatan: memulai, merencanakan, mengerjakan, memeriksa, dan menutup, yang "
    "berjalan bersamaan. Batang biru dan oranye adalah survei Crowe dari slide kuliah: manajer terbaik merencanakan "
    "21 persen. Batang hijau berarsir adalah perkiraan saya untuk Phoenix: merencanakan 8 persen, mengerjakan 86 persen, "
    "memeriksa hanya 3 persen. Angkanya perkiraan, tetapi polanya jelas."))
s.shapes.add_picture(str(HASIL / "Tugas1_DistribusiEffort_Phoenix_Diagram1_Effort_slide.png"),
                     Inches(0.45), Inches(1.35), width=Inches(8.4))
kotak(s, 9.1, 1.45, 3.65, 2.75, fill=BIRU_MUDA)
teks(s, 9.3, 1.58, 3.3, 2.6, [
    [("Lima kelompok kegiatan", {"bold": True, "color": NAVY, "size": 16})],
    "Memulai · merencanakan · mengerjakan · memeriksa · menutup",
    [("Berjalan bersamaan sepanjang proyek, bukan bergiliran.", {"italic": True, "color": ABU})],
], size=15, color=TINTA, space_after=6)
kotak(s, 9.1, 4.35, 3.65, 2.35, fill=MERAH_MUDA)
teks(s, 9.3, 4.48, 3.3, 2.2, [
    [("Phoenix (perkiraan)", {"bold": True, "color": MERAH, "size": 16})],
    [("Merencanakan ", {}), ("~8%", {"bold": True})],
    [("Mengerjakan ", {}), ("~86%", {"bold": True})],
    [("Memeriksa & menguji ", {}), ("~3%", {"bold": True, "color": MERAH})],
], size=16, color=TINTA, space_after=3)

# =============== 5. Tiga ketidakseimbangan ===============
s = slide_baru("Tiga ketidakseimbangan penyebab kegagalan", catatan=(
    "Pertama, saat biaya naik 77 persen, rencana tidak disusun ulang; yang dilakukan membuang fitur. Kedua, pemeriksaan "
    "dan pengujian dikurangi justru saat paling dibutuhkan: fitur yang gagal diuji tidak diuji ulang. Tanggal dan uang "
    "dikunci, jadi mutu yang dikorbankan. Ketiga, tidak ada pemeriksaan sebelum lanjut ke tahap kedua."))
kolom = [
    ("1", "Rencana tidak disusun ulang", "+77%",
     "biaya naik, tetapi yang dilakukan hanya membuang lebih dari 100 fitur", BIRU, BIRU_MUDA),
    ("2", "Pemeriksaan & pengujian dikurangi", "20%",
     "fitur yang gagal diuji tidak diuji ulang; uji coba dibatalkan", MERAH, MERAH_MUDA),
    ("3", "Tak ada pemeriksaan antartahap", "7 → 2 bln",
     "jeda tahap 1 ke tahap 2 dipersingkat; keputusan mulai pakai tidak dicatat", MERAH, MERAH_MUDA),
]
for i, (n, h, big, d, warna, latar) in enumerate(kolom):
    x = 0.6 + i * 4.1
    kotak(s, x, 1.45, 3.9, 4.0, fill=latar)
    lingkaran_nomor(s, x + 0.3, 1.65, n, fill=warna)
    teks(s, x + 0.3, 2.3, 3.3, 0.95, h, size=19, bold=True, color=NAVY)
    teks(s, x + 0.3, 3.2, 3.3, 0.85, big, size=34, bold=True, color=warna, font=JUDUL_FONT)
    teks(s, x + 0.3, 4.1, 3.3, 1.3, d, size=15, color=TINTA)
istilah(s, 0.6, 5.62, 12.15, 1.2, [
    ("Fitur", "kemampuan sistem, misalnya menghitung lembur atau gaji rapel"),
    ("Mengapa mutu dikorbankan", "tanggal dan uang dikunci, jadi yang dikurangi adalah isi pekerjaan dan pengujian (triple constraint)"),
], size=13)

# =============== 6. Waterfall vs Agile ===============
s = slide_baru("Waterfall vs Agile: cara kerja mana yang cocok?", catatan=(
    "Phoenix memakai Waterfall: rencana lengkap, dibangun, lalu diluncurkan sekaligus, sehingga pengujian menumpuk di "
    "akhir dan mudah dikurangi. Agile bekerja dalam putaran pendek yang langsung diuji. Untuk Phoenix yang cocok adalah "
    "gabungan, atau hybrid: bagian dasar direncanakan lengkap, peluncuran bertahap sambil sistem lama tetap jalan."))
segmen = [("Waterfall (berurutan)", BIRU_MUDA, NAVY), ("Hybrid (gabungan)", BIRU, PUTIH), ("Agile (putaran pendek)", BIRU_MUDA, NAVY)]
for i, (lbl, latar, warna_teks) in enumerate(segmen):
    k = kotak(s, 0.6 + i * 4.1, 1.45, 3.9, 0.75, fill=latar)
    tf = k.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = lbl
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = ISI_FONT, Pt(18), True, warna_teks
teks(s, 4.7, 2.25, 3.9, 0.4, "▲ cocok untuk Phoenix", size=14, bold=True, color=BIRU, align=PP_ALIGN.CENTER)
dua = [
    ("Waterfall: menumpuk", ["Rencana besar di awal", "Pengujian menumpuk di akhir", "→ paling mudah dikurangi saat jadwal mepet"]),
    ("Agile: menyebar", ["Rencana kecil di tiap putaran", "Diuji & dicoba pengguna tiap 1–4 minggu", "→ tidak ada satu titik besar untuk dikurangi"]),
]
for i, (h, items) in enumerate(dua):
    x = 0.6 + i * 4.1
    kotak(s, x, 2.8, 3.9, 2.7, fill=BIRU_MUDA)
    teks(s, x + 0.3, 2.95, 3.3, 0.5, h, size=17, bold=True, color=NAVY)
    teks(s, x + 0.3, 3.5, 3.3, 1.95, items, size=15, color=TINTA, space_after=6)
kotak(s, 8.8, 2.8, 3.95, 2.7, fill=PUTIH, line=BIRU)
teks(s, 9.05, 2.95, 3.5, 0.5, "Hybrid untuk Phoenix", size=17, bold=True, color=BIRU)
teks(s, 9.05, 3.5, 3.5, 1.95, [
    "Bagian dasar (aturan gaji, keamanan) direncanakan lengkap",
    "Diluncurkan bertahap per lembaga, sistem lama tetap jalan",
    "Fitur yang belum jelas dikerjakan dalam putaran pendek",
], size=15, color=TINTA, space_after=6)
istilah(s, 0.6, 5.65, 12.15, 1.2, [
    ("Putaran (sprint)", "satu siklus kerja pendek 1–4 minggu"),
    ("Alasan hybrid (PMBOK 7)", "risiko tinggi & aturan ketat → perlu rencana matang; kebutuhan belum pasti & bisa per lembaga → bertahap"),
], size=13)

# =============== 7. Rekomendasi & pelajaran ===============
s = slide_baru("Rekomendasi dan pelajaran utama", catatan=(
    "Empat rekomendasi: buktikan perkiraan manfaat dengan data; bila biaya melenceng jauh, wajib rencana baru dan lapor; "
    "uji coba dan menjalankan sistem lama bersama yang baru tidak boleh dihapus; dan ada pemeriksa dari luar tim yang "
    "lapor ke pimpinan, termasuk keputusan tertulis sebelum sistem dipakai. Pelajarannya: tanggal peluncuran tidak boleh "
    "lebih kuat daripada bukti kesiapan. Pola serupa terjadi di bank TSB dan Coretax, dan relevan untuk BMKG. "
    "Terima kasih."))
rek = [
    ("Memulai", "Buktikan perkiraan manfaat dengan data nyata"),
    ("Merencanakan", "Biaya/jadwal melenceng jauh → wajib rencana baru & lapor"),
    ("Memeriksa", "Uji coba dan sistem lama-baru berjalan bersamaan tidak boleh dihapus"),
    ("Memeriksa & menutup", "Pemeriksa dari luar tim lapor ke pimpinan; keputusan mulai pakai dicatat tertulis"),
]
for i, (pg, r) in enumerate(rek):
    y = 1.45 + i * 1.18
    lingkaran_nomor(s, 0.6, y + 0.1, i + 1)
    teks(s, 1.35, y + 0.02, 6.2, 0.45, pg, size=15, bold=True, color=BIRU)
    teks(s, 1.35, y + 0.42, 6.2, 0.75, r, size=16, color=TINTA)
kotak(s, 7.9, 1.45, 4.85, 4.75, fill=BIRU_MUDA)
teks(s, 8.15, 1.6, 4.35, 0.45, "PELAJARAN UTAMA", size=14, bold=True, color=BIRU)
teks(s, 8.15, 2.05, 4.35, 2.0, "Tanggal peluncuran tidak boleh lebih kuat daripada bukti kesiapan",
     size=24, bold=True, color=NAVY, font=JUDUL_FONT)
teks(s, 8.15, 4.05, 4.35, 2.1, [
    [("Pola serupa: ", {"bold": True}), ("bank TSB, Inggris (2018) dan Coretax DJP (2025).", {})],
    [("Catatan: ", {"bold": True}), ("cara kerja Agile saja tidak cukup; tetap perlu pemeriksa yang tidak memihak.", {})],
    [("Relevan untuk BMKG ", {"bold": True}), ("dan proyek TI pemerintah lain.", {})],
], size=14, color=TINTA, space_after=5)

# =============== 8. Daftar istilah (cadangan) ===============
s = slide_baru("Daftar istilah (cadangan)", catatan="Slide cadangan untuk menjawab pertanyaan; tidak dipresentasikan.")
glos = [
    ("Effort", "tenaga dan waktu proyek"),
    ("Process group", "5 kelompok kegiatan mengelola proyek"),
    ("Initiating", "memulai: memutuskan proyek dan tujuannya"),
    ("Planning", "merencanakan: apa, kapan, biaya, risiko"),
    ("Executing", "mengerjakan: membangun sistem"),
    ("Monitoring & Controlling", "memeriksa, menguji, memperbaiki"),
    ("Closing", "menutup: serah terima, evaluasi"),
    ("PMBOK Guide", "buku standar pengelolaan proyek (PMI)"),
    ("OAG", "lembaga pemeriksa negara Kanada (seperti BPK)"),
    ("Uji coba (pilot)", "dicoba di satu tempat dulu"),
    ("Parallel run", "sistem lama dan baru jalan bersamaan"),
    ("Go-live", "mulai dipakai sungguhan"),
    ("Phase gate", "pemeriksaan sebelum tahap berikutnya"),
    ("Waterfall", "berurutan, diluncurkan sekaligus"),
    ("Agile", "putaran pendek, langsung dicoba pengguna"),
    ("Hybrid", "gabungan Waterfall dan Agile"),
    ("Triple constraint", "3 batasan: isi pekerjaan, waktu, biaya"),
    ("C$", "dolar Kanada"),
]
for k, bagian in enumerate((glos[:9], glos[9:])):
    tbl = s.shapes.add_table(len(bagian), 2, Inches(0.6 + k * 6.2), Inches(1.4), Inches(5.95), Inches(5.2)).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(2.2), Inches(3.75)
    for i, (t, d) in enumerate(bagian):
        for j, v in enumerate((t, d)):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = BIRU_MUDA if i % 2 == 0 else PUTIH
            c.text = ""
            r = c.text_frame.paragraphs[0].add_run(); r.text = v
            r.font.name, r.font.size, r.font.bold = ISI_FONT, Pt(13), (j == 0)
            r.font.color.rgb = NAVY if j == 0 else TINTA

# =============== 9. Daftar pustaka ===============
s = slide_baru("Daftar Pustaka", catatan="Slide cadangan; tidak dipresentasikan.")
teks(s, 0.6, 1.4, 12.1, 5.3, [
    "Crowe, A. (2006). Alpha project managers: What the top 2% know that everyone else does not. Velociteach.",
    "Financial Conduct Authority. (2022). Final notice: TSB Bank plc. https://www.fca.org.uk/publication/final-notices/tsb-bank-plc-2022.pdf",
    "Luthfi, E. T. (2026). IT project management: Pertemuan 4 [Slide perkuliahan]. Program Studi S2 Informatika, Universitas Amikom Yogyakarta.",
    "Office of the Auditor General of Canada. (2018). Report 1—Building and implementing the Phoenix pay system. 2018 Spring Reports of the Auditor General of Canada.",
    "Project Management Institute. (2008). A guide to the project management body of knowledge (PMBOK® Guide) (4th ed.). Project Management Institute.",
    "Project Management Institute. (2021). A guide to the project management body of knowledge (PMBOK® Guide) (7th ed.). Project Management Institute.",
], size=14, color=TINTA, space_after=10)

prs.save(OUT)
print("dibuat:", OUT)
