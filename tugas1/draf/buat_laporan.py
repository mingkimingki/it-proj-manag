"""Membuat laporan Tugas 1 (Analisis Distribusi Effort — Phoenix Pay) dalam format .docx.

Jalankan dengan venv mata kuliah (berisi python-docx):
    ../../.venv/bin/python buat_laporan.py

Hasil: ../hasil/Tugas1_DistribusiEffort_Phoenix_Laporan.docx
Diagram diambil dari ../hasil/*.png (buat dulu dengan buat_diagram.py).

Markup kecil di teks: **tebal**, *miring*, [[disorot kuning]] (penanda yang harus diganti pengguna).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

HASIL = Path(__file__).resolve().parent.parent / "hasil"
OUT = HASIL / "Tugas1_DistribusiEffort_Phoenix_Laporan.docx"
TOKEN = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\]|\*.+?\*)")
C_OAG = "[[CITE: OAG 2018]]"

doc = Document()

# ---------- Pengaturan halaman & gaya ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))
LEBAR = sec.page_width - sec.left_margin - sec.right_margin

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3
for nama, ukuran in (("Heading 1", 14), ("Heading 2", 12), ("Title", 20)):
    st = doc.styles[nama]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(ukuran)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)


def isi_runs(par, teks, ukuran=None):
    for bagian in TOKEN.split(teks):
        if not bagian:
            continue
        if bagian.startswith("**"):
            r = par.add_run(bagian[2:-2]); r.bold = True
        elif bagian.startswith("[["):
            r = par.add_run("[" + bagian[2:-2] + "]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif bagian.startswith("*"):
            r = par.add_run(bagian[1:-1]); r.italic = True
        else:
            r = par.add_run(bagian)
        if ukuran:
            r.font.size = Pt(ukuran)
    return par


def p(teks, align=WD_ALIGN_PARAGRAPH.JUSTIFY, style=None):
    par = doc.add_paragraph(style=style)
    par.alignment = align
    return isi_runs(par, teks)


def h1(teks):
    doc.add_heading(teks, level=1)


def h2(teks):
    isi_runs(doc.add_heading("", level=2), teks)


def bullets(items):
    for it in items:
        par = p(it, style="List Bullet")
        par.paragraph_format.space_after = Pt(3)


def caption(teks):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    isi_runs(par, teks, ukuran=10).italic = True
    for r in par.runs:
        r.italic = True
    par.paragraph_format.space_after = Pt(10)


def gambar(nama, teks_caption, lebar_cm=16):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(HASIL / nama), width=Cm(lebar_cm))
    caption(teks_caption)


def arsir(cell, warna="E8E8E8"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), warna)
    tcPr.append(shd)


def tabel(header, baris, persen, judul=None):
    if judul:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        r = par.add_run(judul); r.italic = True; r.font.size = Pt(10)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    lebar = [int(LEBAR * x / 100) for x in persen]
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        isi_runs(c.paragraphs[0], f"**{h}**", ukuran=10)
        arsir(c)
    for b in baris:
        cells = t.add_row().cells
        for i, v in enumerate(b):
            cells[i].text = ""
            isi_runs(cells[i].paragraphs[0], v, ukuran=10)
    # Lebar kolom harus ditulis juga di <w:tblGrid> (tblLayout "fixed" sudah diset oleh
    # t.autofit = False), kalau tidak Word Online/LibreOffice membagi kolom sama lebar.
    tblPr = t._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(Emu(x).twips for x in lebar))); tblW.set(qn("w:type"), "dxa")
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), lebar):
        gc.set(qn("w:w"), str(Emu(w).twips))
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = lebar[i]
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.1
    # header diulang bila tabel terpotong halaman
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def nomor_halaman(section):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)
    r.font.size = Pt(10)


# ================= SAMPUL =================
for _ in range(6):
    doc.add_paragraph()
t = doc.add_paragraph(style="Title")
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.add_run("Analisis Distribusi Effort Proyek:\nStudi Kasus Phoenix Pay System (Kanada)")
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Tugas 1 — IT Project Management"); r.font.size = Pt(14)
for _ in range(5):
    doc.add_paragraph()
for baris in ("Nama: [[isi nama]]", "NIM: [[isi NIM]]",
              "Program Studi S2 Informatika — Universitas Amikom Yogyakarta",
              "Dosen Pengampu: Emha Taufiq Luthfi, S.T., M.Kom., Ph.D.",
              "Semester Ganjil 2026/2027 — Oktober 2026"):
    p(baris, align=WD_ALIGN_PARAGRAPH.CENTER)

# Isi di section baru (sampul tanpa nomor halaman)
isi = doc.add_section(WD_SECTION.NEW_PAGE)
isi.footer.is_linked_to_previous = False
nomor_halaman(isi)

# ================= 1 =================
h1("1. Pendahuluan")
p("Laporan ini menganalisis bagaimana tenaga dan waktu (*effort*) sebuah proyek teknologi informasi (TI) nyata terbagi ke dalam lima kelompok proses manajemen proyek (*process group*), lalu menilai apakah pembagian itu ikut menentukan kegagalan proyek. Kasus yang dipilih adalah **Phoenix Pay System**, sistem penggajian (*payroll*, yaitu sistem yang menghitung dan membayarkan gaji pegawai) milik Pemerintah Kanada yang mulai dipakai pada 2016.")
p(f"Kasus ini dipilih karena terdokumentasi secara resmi dan independen oleh Office of the Auditor General of Canada (OAG), lembaga pemeriksa keuangan negara Kanada yang setara dengan BPK, dalam laporan audit tahun 2018 {C_OAG}. OAG memakai *PMBOK Guide*, yaitu buku panduan standar manajemen proyek dari Project Management Institute (PMI), sebagai salah satu kriteria auditnya, sehingga kerangkanya sejalan dengan materi kuliah. Laporan tersebut juga memuat angka konkret tentang pemangkasan lingkup, pengujian, jadwal, dan staf, yang menjadi bahan utama analisis *effort*.")

# ================= 2 =================
h1("2. Landasan Teori")
h2("2.1 Lima *process group*")
p("PMBOK edisi ke-4 mengelompokkan kegiatan manajemen proyek ke dalam lima *process group* [[CITE: PMI 2008]]:")
bullets([
    "**Initiating** (inisiasi): mengesahkan dimulainya proyek, tujuan, dan pihak yang berkepentingan (*stakeholder*, yaitu orang atau organisasi yang terdampak proyek).",
    "**Planning** (perencanaan): menyusun rencana lingkup, jadwal, biaya, mutu, dan risiko.",
    "**Executing** (pelaksanaan): mengerjakan pekerjaan sesuai rencana.",
    "**Monitoring & Controlling** (pemantauan dan pengendalian): memantau kemajuan, menguji hasil, serta mengendalikan perubahan dan risiko.",
    "**Closing** (penutupan): menyerahkan hasil, menutup fase atau proyek, dan mencatat pelajaran.",
])
p("Kelima kelompok ini bukan tahap berurutan, melainkan tumpang tindih sepanjang proyek; *Monitoring & Controlling* berlangsung dari awal hingga akhir [[CITE: PMI 2008; Luthfi 2026]]. PMBOK juga menunjukkan bahwa biaya memperbaiki kesalahan makin mahal menjelang akhir proyek.")
h2("2.2 Patokan distribusi waktu")
p("Materi kuliah menyajikan data survei Crowe tentang porsi waktu manajer proyek per *process group* [[CITE: Crowe 2006]]. Manajer proyek “Alpha” (sekitar 2% terbaik menurut penilaian *stakeholder*) memakai 21% waktunya untuk *Planning* dan 69% untuk *Executing*, sedangkan manajer proyek lain 11% dan 82%. Data ini korelasional dan mengukur waktu manajer proyek, bukan seluruh tim, sehingga dalam laporan ini dipakai sebagai **pembanding pola**, bukan angka pasti.")
h2("2.3 Pendekatan pengembangan")
p("PMBOK edisi ke-7 memandang cara mengerjakan proyek sebagai spektrum [[CITE: PMI 2021]]: **predictive** atau Waterfall (tahap berurutan dengan rencana lengkap di awal), **adaptive** atau Agile (siklus pendek 1–4 minggu dengan umpan balik pengguna di setiap siklus), dan **hybrid** (campuran keduanya). Pilihannya ditentukan oleh faktor seperti kepastian kebutuhan, kemungkinan menyerahkan hasil secara bertahap, risiko, regulasi, serta ukuran dan budaya organisasi.")

# ================= 3 =================
h1("3. Studi Kasus: Phoenix Pay System")
p(f"Pada 2009, Pemerintah Kanada memulai penggantian sistem gaji berusia 40 tahun untuk sekitar 290.000 pegawai di 101 instansi, dengan target penghematan sekitar C$70 juta per tahun (C$ = dolar Kanada). Kementerian Public Services and Procurement Canada (PSPC) memimpin proyek dan pada 2011 mengontrak IBM untuk mengkustomisasi (mengubah agar sesuai kebutuhan) *software* (perangkat lunak) penggajian PeopleSoft. Sistem diluncurkan dalam dua gelombang: 34 instansi pada 24 Februari 2016 dan 67 instansi pada 21 April 2016 {C_OAG}. Lini waktu proyek ditunjukkan pada Gambar 1.")
gambar("Tugas1_DistribusiEffort_Phoenix_Diagram2_LiniWaktu.png",
       "Gambar 1. Lini waktu proyek Phoenix Pay dan titik kendali yang dilewati (diolah dari OAG, 2018)")
p(f"Hasilnya buruk. Pada Juni 2017, nilai gaji yang salah atau tertunda melebihi C$520 juta, sekitar 51% pegawai menerima slip gaji yang salah, dan lebih dari 494.500 permintaan perbaikan menumpuk. OAG menyebut proyek ini sebagai kegagalan manajemen proyek dan pengawasan yang “*incomprehensible*” {C_OAG}. Temuan utamanya:")
bullets([
    "Ketika perkiraan biaya IBM (C$274 juta) hampir dua kali anggaran (C$155 juta), lebih dari 100 fungsi gaji dihapus atau ditunda, pengujian dikurangi, jarak antargelombang dipadatkan dari 7 menjadi 2 bulan, dan staf dikurangi, tanpa perencanaan ulang.",
    "Sistem tidak diuji secara utuh; 20% fungsi yang diperiksa auditor gagal uji dan tidak diuji ulang; uji coba (*pilot*, yaitu menjalankan sistem di lingkungan terbatas lebih dulu) dibatalkan pada Juni 2015.",
    "Saran konsultan Gartner untuk meluncurkan bertahap dan menjalankan sistem lama dan baru bersamaan (*parallel run*) tidak dipertimbangkan.",
    "Tidak ada pengawasan independen: informasi ke wakil menteri hanya datang dari tim proyek, dan audit internal yang direncanakan empat kali tidak pernah dilakukan.",
    "Keputusan *go-live* (mulai memakai sistem secara resmi) tidak didokumentasikan.",
])

# ================= 4 =================
h1("4. Analisis Distribusi Effort")
p("Tidak ada data jam kerja per *process group*, sehingga distribusi *effort* diestimasi dengan memetakan temuan OAG ke *process group* menurut Tabel 3-1 PMBOK 4 (misalnya, pengujian termasuk *Monitoring & Controlling*), lalu membandingkan polanya dengan data Crowe. Sudut pandang yang dipakai sama dengan Crowe, yaitu porsi perhatian manajemen proyek.")
tabel(["Process group", "PM Alpha", "PM lain", "Phoenix (estimasi)", "Alasan utama"], [
    ["Initiating", "2%", "1%", "2%", "*Business case* (dokumen alasan, biaya, dan manfaat proyek) ada, tetapi asumsi penghematan tidak diuji dan instansi pengguna tidak dilibatkan."],
    ["Planning", "21%", "11%", "~8%", "Tidak ada perencanaan ulang saat biaya naik 77%; rencana kontinjensi (rencana cadangan) dibuat kurang dari dua minggu sebelum peluncuran."],
    ["Executing", "69%", "82%", "~86%", "*Effort* terpusat pada membangun sistem dan memindahkan data demi menjaga jadwal dan anggaran."],
    ["Monitoring & Controlling", "5%", "4%", "~3%", "*Pilot* dibatalkan, pengujian tidak utuh, tanpa *parallel run*, audit internal, atau tinjauan kesiapan yang independen."],
    ["Closing", "3%", "2%", "~1%", "Tidak ada evaluasi berarti antargelombang; keputusan *go-live* tidak dicatat."],
], [19, 10, 10, 13, 48], judul="Tabel 1. Estimasi distribusi effort Phoenix dibandingkan patokan Crowe (2006)")
gambar("Tugas1_DistribusiEffort_Phoenix_Diagram1_Effort.png",
       "Gambar 2. Distribusi effort per process group: Phoenix (estimasi) dibandingkan data Crowe (2006)")
p("Angka Phoenix adalah estimasi penulis, bukan data terukur; yang dapat dipertanggungjawabkan adalah polanya. Bila setiap angka digeser ±3 poin, kesimpulannya tetap sama. Distribusi ini timpang dalam tiga hal.")
p("**Pertama, Planning terlalu kecil dan hanya di awal.** PMBOK menyatakan bahwa *process group* berulang, sehingga kenaikan biaya 77% seharusnya memicu perencanaan ulang. Phoenix justru hanya memangkas lingkup, pola yang mirip, bahkan lebih rendah daripada, “manajer proyek lain” dalam data Crowe.")
p("**Kedua, Monitoring & Controlling dipangkas ketika paling dibutuhkan.** Dari *triple constraint* (tiga kendala proyek: lingkup, waktu, biaya), waktu dan biaya dikunci, sehingga lingkup dan mutu yang dikorbankan. Padahal mutu dijaga oleh pengujian, uji coba, dan pengawasan. Karena biaya perbaikan kesalahan paling mahal setelah sistem dipakai, penghematan di tahap ini dibayar mahal: lebih dari C$520 juta gaji salah atau tertunda dalam setahun.")
p("**Ketiga, akhir fase tidak berfungsi sebagai** *phase gate* (titik evaluasi untuk memutuskan lanjut, menunda, atau berhenti). Memadatkan jarak antargelombang dari 7 menjadi 2 bulan menghapus kesempatan belajar dari gelombang pertama.")
p("Jadi, ketimpangan *effort*, yaitu *Executing* yang terlalu dominan, menjadi faktor utama kegagalan Phoenix. Namun yang lebih menentukan adalah **waktu dan mutu** *effort* itu: *Planning* tidak diulang saat dibutuhkan, dan *Monitoring & Controlling* dipangkas tepat sebelum peluncuran.")

# ================= 5 =================
h1("5. Perbandingan Waterfall dan Agile")
p("Phoenix dikerjakan secara *predictive* (Waterfall): anggaran dan lingkup ditetapkan di awal, sistem dibangun lengkap, lalu diluncurkan ke 101 instansi dalam dua gelombang besar. Dalam pendekatan ini *Planning* menumpuk di awal dan *Monitoring & Controlling*, terutama pengujian, menumpuk menjelang peluncuran, sehingga menjadi bagian yang paling mudah dipangkas ketika jadwal terdesak. Dalam pendekatan *adaptive* (Agile), perencanaan dilakukan sedikit demi sedikit di setiap siklus, serta pengujian dan demo kepada pengguna terjadi setiap 1–4 minggu, sehingga tidak ada satu titik besar yang bisa dipangkas sekaligus (Tabel 2).")
tabel(["Process group", "Waterfall (predictive)", "Agile (adaptive)"], [
    ["Planning", "Besar di awal; diulang hanya bila ada perubahan besar", "Kecil di awal, lalu di setiap siklus; daftar kebutuhan (*backlog*) diprioritaskan ulang terus"],
    ["Executing", "Besar di tengah proyek", "Tersebar di setiap siklus"],
    ["Monitoring & Controlling", "Memuncak saat pengujian menjelang peluncuran", "Di setiap siklus: pengujian dan demo ke pengguna"],
    ["Closing", "Di akhir fase dan akhir proyek", "Evaluasi kecil tiap siklus (*retrospective*) dan penutupan akhir"],
], [24, 38, 38], judul="Tabel 2. Perbedaan distribusi effort pada Waterfall dan Agile")
p("Faktor-faktor PMBOK 7 tidak menunjuk ke satu ujung spektrum. Risiko yang sangat tinggi, regulasi pemerintah, organisasi yang hierarkis, dan tim yang besar mengarah ke *predictive*. Sebaliknya, kebutuhan yang belum pasti karena instansi pengguna tidak dilibatkan, serta kemungkinan meluncurkan sistem per instansi, mengarah ke pendekatan bertahap dan *adaptive*.")
p("Karena itu, pendekatan yang paling tepat adalah **hybrid**: fondasi (aturan gaji inti, integrasi, keamanan, dan privasi) direncanakan secara *predictive* bersama instansi pengguna; peluncuran dilakukan bertahap (*incremental*) mulai dari instansi dengan aturan gaji sederhana, disertai *parallel run* dan *phase gate* sebelum setiap gelombang; fungsi yang paling tidak pasti, seperti gaji rapel (selisih gaji yang dibayar mundur), dikembangkan dalam siklus pendek bersama petugas gaji. Pilihan ini sejalan dengan saran Gartner yang diabaikan.")
p("Perlu dicatat, Agile bukan obat untuk semua masalah Phoenix. Masalah tata kelola (*governance*, yaitu siapa yang mengawasi dan berwenang memutuskan), seperti informasi yang hanya datang dari tim proyek dan keputusan *go-live* yang tidak dicatat, tetap bisa terjadi pada proyek Agile. Pendekatan apa pun perlu dibarengi pengawasan independen.")

# ================= 6 =================
h1("6. Rekomendasi")
tabel(["No", "Process group", "Rekomendasi"], [
    ["1", "Initiating", "Uji asumsi manfaat dengan data nyata sebelum proyek disahkan."],
    ["2", "Initiating", "Libatkan instansi pengguna dalam komite pengarah (*steering committee*) sejak awal."],
    ["3", "Planning", "Wajibkan perencanaan ulang dan pelaporan ke pemberi anggaran bila biaya atau jadwal menyimpang melebihi batas tertentu (misalnya 10–20%), bukan memangkas lingkup diam-diam."],
    ["4", "Planning", "Siapkan dan uji rencana kontinjensi berbulan-bulan sebelum peluncuran."],
    ["5", "Executing", "Latih pengguna dengan sistem yang sesungguhnya."],
    ["6", "Monitoring & Controlling", "Jadikan pengujian utuh, *pilot*, dan *parallel run* syarat yang tidak boleh dipangkas tanpa persetujuan pihak independen."],
    ["7", "Monitoring & Controlling", "Adakan audit internal dan tinjauan kesiapan independen yang melapor langsung ke pimpinan."],
    ["8", "Closing", "Catat keputusan *go/no-go* dengan kriteria dan nama pengambil keputusan; beri masa stabilisasi dan evaluasi sebelum gelombang berikutnya."],
], [7, 25, 68], judul="Tabel 3. Rekomendasi per process group")
p("Untuk distribusi *effort*, porsi *Planning* dan *Monitoring & Controlling* sebaiknya diberi alokasi anggaran dan jadwal tersendiri yang tidak boleh dipakai menutup kekurangan *Executing*, dengan pola manajer proyek Alpha sebagai acuan.")
p("Pola yang sama terlihat pada migrasi sistem TSB Bank di Inggris pada 2018, ketika jadwal disusun mundur dari tanggal target dan pengujian dipadatkan [[CITE: FCA 2022]], serta pada peluncuran Coretax DJP pada Januari 2025, ketika kendala *login* dan penerbitan faktur muncul saat banyak wajib pajak bertransaksi serentak [[CITE: Luthfi 2026]]. Bagi proyek TI pemerintah di Indonesia, termasuk peralihan sistem observasi di BMKG, pelajarannya sama: **tanggal peluncuran tidak boleh lebih kuat daripada bukti kesiapan.**")

# ================= 7 =================
h1("7. Kesimpulan")
p("Kegagalan Phoenix menunjukkan bahwa distribusi *effort* yang terlalu berat pada *Executing*, dengan *Planning* yang tidak diulang dan *Monitoring & Controlling* yang dipangkas, berujung pada sistem yang diluncurkan sebelum siap. Pendekatan *hybrid* dengan peluncuran bertahap dan pengawasan independen lebih tepat untuk proyek berskala nasional seperti ini.")

# ================= Glosarium =================
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
h1("Glosarium")
glos = [
    ("Adaptive / Agile", "Pendekatan proyek dalam siklus pendek dengan umpan balik pengguna di setiap siklus"),
    ("Audit internal", "Pemeriksaan oleh unit pemeriksa di dalam organisasi yang melapor langsung ke pimpinan, terpisah dari tim proyek"),
    ("Backlog", "Daftar kebutuhan atau fitur yang diurutkan berdasarkan prioritas dalam pendekatan Agile"),
    ("Business case", "Dokumen yang menjelaskan alasan, biaya, dan manfaat sebuah proyek"),
    ("Effort", "Jumlah tenaga dan waktu yang dicurahkan untuk suatu pekerjaan"),
    ("Gaji rapel", "Selisih gaji yang dibayarkan mundur, misalnya setelah kenaikan pangkat"),
    ("Go-live", "Saat sistem baru mulai dipakai secara resmi"),
    ("Go/no-go", "Keputusan resmi apakah sistem boleh mulai dipakai (go) atau ditunda (no-go)"),
    ("Hybrid", "Campuran pendekatan predictive dan adaptive"),
    ("Incremental (bertahap)", "Hasil proyek diserahkan sedikit demi sedikit; setiap tahap menambah fungsi atau pengguna"),
    ("Kustomisasi", "Mengubah software siap pakai agar sesuai kebutuhan khusus"),
    ("OAG", "Office of the Auditor General of Canada; lembaga pemeriksa keuangan dan kinerja pemerintah Kanada, setara BPK"),
    ("Parallel run", "Menjalankan sistem lama dan baru bersamaan untuk sementara dan membandingkan hasilnya"),
    ("Payroll", "Sistem yang menghitung dan membayarkan gaji pegawai"),
    ("Phase gate", "Titik evaluasi di akhir fase untuk memutuskan proyek lanjut, ditunda, atau dihentikan"),
    ("Pilot (uji coba)", "Menjalankan sistem baru di lingkungan terbatas sebelum dipakai seluruhnya"),
    ("PMBOK Guide", "Buku panduan standar manajemen proyek dari Project Management Institute (PMI)"),
    ("Predictive / Waterfall", "Pendekatan proyek berurutan dengan rencana lengkap di awal"),
    ("Process group", "Kelompok proses manajemen proyek: Initiating, Planning, Executing, Monitoring & Controlling, Closing"),
    ("Rencana kontinjensi", "Rencana cadangan bila sistem bermasalah"),
    ("Retrospective", "Evaluasi cara kerja tim di akhir setiap siklus Agile"),
    ("Software", "Perangkat lunak, yaitu program komputer"),
    ("Stakeholder", "Orang atau organisasi yang terdampak atau berkepentingan dengan proyek"),
    ("Steering committee", "Komite pimpinan yang mengarahkan dan mengawasi proyek"),
    ("Tata kelola (governance)", "Pengaturan siapa yang mengawasi proyek dan siapa yang berwenang memutuskan"),
    ("Triple constraint", "Tiga kendala utama proyek yang saling memengaruhi: lingkup, waktu, dan biaya"),
]
tabel(["Istilah", "Penjelasan"], [[a, b] for a, b in glos], [30, 70])

# ================= Daftar pustaka =================
h1("Daftar Pustaka")
p("[[Sisipkan daftar pustaka: Mendeley Cite → Insert Bibliography (APA 7th edition). Setelah itu hapus kalimat ini dan daftar acuan di bawah.]]", align=WD_ALIGN_PARAGRAPH.LEFT)
p("*Daftar acuan (untuk dimasukkan ke Mendeley):*", align=WD_ALIGN_PARAGRAPH.LEFT)
for ref in [
    "Crowe, A. (2006). *Alpha project managers: What the top 2% know that everyone else does not*. Velociteach.",
    "Financial Conduct Authority. (2022). *Final notice: TSB Bank plc*. https://www.fca.org.uk/publication/final-notices/tsb-bank-plc-2022.pdf",
    "Luthfi, E. T. (2026). *IT project management: Pertemuan 4* [Slide perkuliahan]. Program Studi S2 Informatika, Universitas Amikom Yogyakarta.",
    "Office of the Auditor General of Canada. (2018). *Report 1—Building and implementing the Phoenix pay system*. 2018 Spring Reports of the Auditor General of Canada. https://www.oag-bvg.gc.ca/internet/English/parl_oag_201805_01_e_43033.html",
    "Project Management Institute. (2008). *A guide to the project management body of knowledge (PMBOK® Guide)* (4th ed.). Project Management Institute.",
    "Project Management Institute. (2021). *A guide to the project management body of knowledge (PMBOK® Guide)* (7th ed.). Project Management Institute.",
]:
    par = p(ref, align=WD_ALIGN_PARAGRAPH.LEFT)
    par.paragraph_format.left_indent = Cm(1.27)
    par.paragraph_format.first_line_indent = Cm(-1.27)

# Templat bawaan python-docx punya <w:zoom> tanpa atribut wajib "percent" — lengkapi agar lolos validasi.
for zoom in doc.settings.element.findall(qn("w:zoom")):
    if zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

doc.save(OUT)
print("dibuat:", OUT)
