"""Membuat PDF bukti ClickUp (pengganti "file tools digital", karena fitur Export ClickUp berbayar).

Jalankan dengan venv mata kuliah (berisi python-docx dan Pillow):
    ../../.venv/bin/python buat_bukti_pdf.py

Langkah:
    1. Menyatukan bukti/01_gantt_atas.png dan bukti/02_gantt_bawah.png menjadi bukti/03_gantt_utuh.png.
       Kedua tangkapan layar diambil pada posisi geser horizontal yang sama; gambar bawah adalah gambar atas
       yang digulir GESER piksel. Sambungan diletakkan di batas baris 1.5.5 / 1.5.6.
    2. Membuat ../hasil/Tugas2_Bukti_ClickUp.pdf: halaman keterangan (A4 mendatar) + Gantt utuh
       (ukuran halaman mengikuti gambar, supaya bisa diperbesar di layar). PDF dibuat dengan LibreOffice.
    3. Menyalin clickup_import_spdc.csv ke ../hasil/Tugas2_ClickUp_Import.csv.
Bila tangkapan layar diganti, periksa ulang GESER dan SAMBUNG.
"""
import datetime as dt
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageChops

from siapkan_clickup import NAMA_LIST, data

DIR = Path(__file__).resolve().parent
BUKTI = DIR / "bukti"
HASIL = DIR.parent / "hasil"
ATAS, BAWAH, UTUH = BUKTI / "01_gantt_atas.png", BUKTI / "02_gantt_bawah.png", BUKTI / "03_gantt_utuh.png"
PDF = HASIL / "Tugas2_Bukti_ClickUp.pdf"
CSV = HASIL / "Tugas2_ClickUp_Import.csv"
GESER = 666     # baris y di gambar bawah = baris y + GESER di gambar atas
SAMBUNG = 891   # baris potong di gambar atas (batas antara baris 1.5.5 dan 1.5.6)
LEBAR_GAMBAR = Cm(50)   # lebar Gantt utuh di PDF; batas ukuran halaman Word ±55 cm
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober",
         "November", "Desember"]


def satukan():
    a, b = Image.open(ATAS).convert("RGB"), Image.open(BAWAH).convert("RGB")
    # cek tumpang-tindih: 40 baris di atas sambungan harus hampir sama. Selisih kecil wajar: garis ketergantungan
    # tugas yang tersorot kursor saat gambar bawah diambil tampil lebih gelap.
    per_baris = sorted(sum(ImageChops.difference(a.crop((45, y, 1890, y + 1)), b.crop((45, y - GESER, 1890, y - GESER + 1)))
                           .convert("L").histogram()[30:]) for y in range(SAMBUNG - 40, SAMBUNG))
    if per_baris[len(per_baris) // 2] > 20:
        raise SystemExit(f"Gambar atas dan bawah tidak cocok di sambungan (median {per_baris[len(per_baris) // 2]} "
                         "piksel berbeda per baris); periksa GESER/SAMBUNG.")
    utuh = Image.new("RGB", (a.width, SAMBUNG + b.height - (SAMBUNG - GESER)), "white")
    utuh.paste(a.crop((0, 0, a.width, SAMBUNG)), (0, 0))
    utuh.paste(b.crop((0, SAMBUNG - GESER, b.width, b.height)), (0, SAMBUNG))
    utuh.save(UTUH)
    print("Tersimpan:", UTUH, utuh.size)
    return utuh.size


def tanggal(t):
    return f"{t.day} {BULAN[t.month - 1]} {t.year}"


def p(doc, teks, ukuran=12, tebal=False, align=WD_ALIGN_PARAGRAPH.LEFT, warna=None, sesudah=6):
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.space_after = Pt(sesudah)
    r = par.add_run(teks)
    r.font.size, r.bold = Pt(ukuran), tebal
    if warna:
        r.font.color.rgb = warna
    return par


def buat_pdf(ukuran_gambar):
    fase, tugas, milestone = data()
    kritis = [t for t in tugas if "jalur-kritis" in t["tag"]]
    m_kritis = [m["kode"] for m in milestone if "jalur-kritis" in m["tag"]]
    n_ketergantungan = sum(len(x["pendahulu"]) for x in tugas + milestone)
    mulai, selesai = min(t["mulai"] for t in tugas), max(m["selesai"] for m in milestone)
    diambil = dt.date.fromtimestamp(ATAS.stat().st_mtime)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))

    # ---------- Halaman 1: keterangan ----------
    p(doc, "Bukti File Tools Digital: ClickUp", 20, True, sesudah=2)
    p(doc, "Rencana Proyek Sistem Peringatan Dini Cuaca Publik (SPDC) — Tugas Kelompok IT Project Management",
      13, warna=RGBColor(0x40, 0x40, 0x40), sesudah=14)
    baris = [
        ("Alat", "ClickUp, paket Free; tampilan Gantt dengan fitur jalur kritis (Critical Path) dinyalakan"),
        ("Lokasi data", f"Space \"Tugas 2 ITPM\" → List \"{NAMA_LIST}\" (level 1 WBS)"),
        ("Isi", f"{len(fase)} fase sebagai tugas (level 2), {len(tugas)} work package sebagai subtugas (level 3), "
                f"{len(milestone)} milestone, dan {n_ketergantungan} ketergantungan Finish-to-Start"),
        ("Jadwal", f"{tanggal(mulai)} sampai {tanggal(selesai)}; hari kerja Senin–Jumat (hari libur nasional tidak dihitung)"),
        ("Jalur kritis", f"Dihitung sendiri oleh ClickUp (batang merah) dan sama dengan perhitungan CPM di laporan "
                         f"(Lampiran B): {len(kritis)} work package dengan float 0 serta milestone {', '.join(m_kritis)}"),
        ("Cara pengisian", "Seluruh data diambil dari laporan dan dikirim ke ClickUp lewat API (skrip buat_clickup.py), "
                           "sehingga isinya sama persis dengan laporan"),
        ("Berkas data", "Tugas2_ClickUp_Import.csv: tugas, subtugas, deskripsi, tanggal, dan tag; dapat diimpor ulang "
                        "ke ClickUp (Import → Spreadsheet/CSV). Ketergantungan tidak ikut lewat CSV"),
        ("Catatan", "Fitur Export ClickUp hanya tersedia pada paket berbayar. Karena itu bukti diberikan berupa "
                    "tangkapan layar (halaman berikut) dan berkas CSV"),
        ("Tangkapan layar", f"Diambil {tanggal(diambil)}; dua tangkapan layar disatukan menjadi satu gambar utuh. "
                            "Perbesar halaman berikut untuk membaca nama tugas"),
    ]
    t = doc.add_table(rows=len(baris), cols=2)
    t.style = "Table Grid"
    t.autofit = False
    lebar = (Cm(5), Cm(20.7))
    # LibreOffice memakai <w:tblGrid>; tanpa ini kolom dibagi sama lebar
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), lebar):
        gc.set(qn("w:w"), str(w.twips))
    for i, (k, v) in enumerate(baris):
        for j, (isi, tebal) in enumerate(((k, True), (v, False))):
            c = t.cell(i, j)
            c.width = lebar[j]
            r = c.paragraphs[0].add_run(isi)
            r.font.size, r.bold = Pt(11.5), tebal
            c.paragraphs[0].paragraph_format.space_after = Pt(2)

    # ---------- Halaman 2: Gantt utuh, halaman seukuran gambar ----------
    lebar_px, tinggi_px = ukuran_gambar
    tinggi_gambar = int(LEBAR_GAMBAR * tinggi_px / lebar_px)
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = LEBAR_GAMBAR + Cm(2), tinggi_gambar + Cm(3.2)
    for side in ("left_margin", "right_margin", "bottom_margin"):
        setattr(s, side, Cm(1))
    s.top_margin = Cm(1)
    p(doc, "Gantt Chart lengkap di ClickUp (batang merah = jalur kritis; garis = ketergantungan; ◆ = milestone)",
      14, True, sesudah=4)
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.add_run().add_picture(str(UTUH), width=LEBAR_GAMBAR)

    with tempfile.TemporaryDirectory() as tmp:
        docx = Path(tmp) / (PDF.stem + ".docx")
        doc.save(docx)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HASIL), str(docx)],
                       check=True, capture_output=True)
    print("Tersimpan:", PDF)


if __name__ == "__main__":
    HASIL.mkdir(exist_ok=True)
    buat_pdf(satukan())
    shutil.copyfile(DIR / "clickup_import_spdc.csv", CSV)
    print("Tersimpan:", CSV)
