"""Membuat slide presentasi Tugas 2 (Rencana Proyek SPDC), 10 menit, 3 penyaji.

11 slide dipresentasikan + 3 slide cadangan (jalur kritis, daftar istilah, daftar pustaka).
Penyaji: Eka (slide 1–5, pembuka dan WBS), Erlisnawati (6–8, jadwal), Sri Rejeki (9–11, risiko dan penutup).
Naskah tiap slide ada di catatan pembicara (speaker notes).

Jalankan dengan venv mata kuliah (berisi python-pptx):
    ../../.venv/bin/python buat_slide.py

Hasil: ../hasil/Tugas2_RencanaProyek_SPDC_Slide.pptx dan .pdf (PDF dibuat dengan LibreOffice).
Angka diambil dari laporan_tugas2.md (tabel jadwal, kamus WBS, Risk Register), jadi selalu sama dengan laporan.
Butuh: ../clickup/bukti/02_gantt_bawah.png dan ../../tugas1/sumber/logo_amikom.png.
"""
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

DRAF = Path(__file__).resolve().parent
TUGAS = DRAF.parent
HASIL = TUGAS / "hasil"
LOGO = TUGAS.parent / "tugas1" / "sumber" / "logo_amikom.png"
GANTT = TUGAS / "clickup" / "bukti" / "02_gantt_bawah.png"
OUT = HASIL / "Tugas2_RencanaProyek_SPDC_Slide.pptx"

sys.path.insert(0, str(TUGAS / "clickup"))
from siapkan_clickup import hari_kerja  # noqa: E402  (hari kerja ke-n → tanggal kalender, sama dengan ClickUp)

PUTIH = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TINTA = RGBColor(0x1F, 0x1F, 0x1F)
ABU = RGBColor(0x52, 0x51, 0x4E)
REDUP = RGBColor(0x89, 0x87, 0x81)
BIRU = RGBColor(0x2A, 0x78, 0xD6)
BIRU_MUDA = RGBColor(0xEA, 0xF2, 0xFC)
MERAH = RGBColor(0xD0, 0x3B, 0x3B)
MERAH_MUDA = RGBColor(0xFB, 0xEA, 0xEA)
ABU_MUDA = RGBColor(0xF3, 0xF3, 0xF1)
WARNA_LEVEL = {"Rendah": RGBColor(0xC6, 0xE0, 0xB4), "Sedang": RGBColor(0xFF, 0xE6, 0x99),
               "Tinggi": RGBColor(0xF4, 0xB1, 0x83), "Kritikal": RGBColor(0xFF, 0x7C, 0x80)}
JUDUL_FONT, ISI_FONT = "Cambria", "Calibri"
BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]

# ---------- Data dari laporan ----------
teks_md = (DRAF / "laporan_tugas2.md").read_text(encoding="utf8")


def sel(baris):
    return [c.strip() for c in baris.strip().strip("|").split("|")]


def bersih(c):
    return c.replace("*", "").strip()


def ambil_tabel(kepala):
    blok = []
    for b in teks_md[teks_md.index(kepala):].splitlines():
        if not b.startswith("|"):
            break
        blok.append(sel(b))
    return blok[2:]


AKTIVITAS = {bersih(c[0]): dict(nama=bersih(c[1]), dur=int(bersih(c[6])), es=int(c[8]), ef=int(c[9]),
                                fl=int(bersih(c[12]))) for c in ambil_tabel("| Kode | Aktivitas | O |")}
TOTAL = max(a["ef"] for a in AKTIVITAS.values())
FASE = []   # [(kode, nama, [(kode_wp, nama_wp), ...]), ...]
for c in ambil_tabel("| Kode | Work package | Uraian |"):
    kode, nama = bersih(c[0]), bersih(c[1])
    if kode.count(".") == 1:
        FASE.append((kode, nama, []))
    else:
        FASE[-1][2].append((kode, nama))
RISIKO = []
for c in ambil_tabel("| ID | Kategori"):
    k2, d2 = (int(x) for x in c[10].split(" = ")[0].split(" × "))
    RISIKO.append(dict(id=bersih(c[0]), k=int(c[4]), d=int(c[5]), k2=k2, d2=d2))
JUMLAH_WP = sum(len(f[2]) for f in FASE)
MILESTONE = [k for k in AKTIVITAS if k.startswith("M")]


def tanggal(hari):
    d = hari_kerja(hari)
    return f"{d.day} {BULAN[d.month - 1]}"


def level(nilai):
    return "Kritikal" if nilai >= 20 else "Tinggi" if nilai >= 13 else "Sedang" if nilai >= 6 else "Rendah"


# ---------- Pembantu gambar slide (pola sama dengan Tugas 1) ----------
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
KOSONG = prs.slide_layouts[6]
nomor = {"n": 0}


def teks(slide, x, y, w, h, isi, size=16, color=TINTA, bold=False, italic=False,
         font=ISI_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    """isi: str, atau list paragraf; tiap paragraf str atau list run (teks, {opsi})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.05))
    paragraf = isi if isinstance(isi, list) else [isi]
    for i, par in enumerate(paragraf):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = par if isinstance(par, list) else [(par, {})]
        for t, opt in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", italic)
            f.color.rgb = opt.get("color", color)
    return tb


def kotak(slide, x, y, w, h, fill=BIRU_MUDA, bentuk=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08):
    s = slide.shapes.add_shape(bentuk, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if bentuk == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def label(shape, isi, size=14, color=TINTA, bold=False, align=PP_ALIGN.CENTER):
    """Teks di dalam bentuk (kotak, panah); isi: str atau list baris."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.08))
    for i, baris_ in enumerate(isi if isinstance(isi, list) else [isi]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = baris_
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = ISI_FONT, Pt(size), bold, color


def panah(slide, x, y, w=0.3, h=0.3, arah=MSO_SHAPE.RIGHT_ARROW, fill=REDUP):
    return kotak(slide, x, y, w, h, fill=fill, bentuk=arah)


def garis(slide, x, y, w, h, fill=REDUP):
    return kotak(slide, x, y, w, h, fill=fill, bentuk=MSO_SHAPE.RECTANGLE)


def tabel_sel(cell, isi, size=12, bold=False, color=TINTA, fill=None, align=PP_ALIGN.LEFT):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.text = ""
    cell.margin_left = cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, baris_ in enumerate(isi if isinstance(isi, list) else [isi]):
        p = cell.text_frame.paragraphs[0] if i == 0 else cell.text_frame.add_paragraph()
        p.alignment = align
        runs = baris_ if isinstance(baris_, list) else [(baris_, {})]
        for t, opt in runs:
            r = p.add_run()
            r.text = t
            r.font.name = ISI_FONT
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", bold)
            r.font.color.rgb = opt.get("color", color)


def slide_baru(judul=None, catatan="", penyaji=None, judul_slide=False):
    s = prs.slides.add_slide(KOSONG)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PUTIH
    nomor["n"] += 1
    if judul:
        teks(s, 0.6, 0.35, 12.1, 0.9, judul, size=30, bold=True, color=NAVY, font=JUDUL_FONT,
             anchor=MSO_ANCHOR.MIDDLE)
    if judul_slide:
        s.shapes.add_picture(str(LOGO), Inches(11.55), Inches(0.45), height=Inches(1.1))
    else:
        s.shapes.add_picture(str(LOGO), Inches(0.35), Inches(6.98), height=Inches(0.4))
        teks(s, 12.2, 7.0, 0.8, 0.35, str(nomor["n"]), size=11, color=REDUP, align=PP_ALIGN.RIGHT)
        if penyaji:
            teks(s, 8.0, 7.0, 4.2, 0.35, f"Penyaji: {penyaji}", size=11, color=REDUP, align=PP_ALIGN.RIGHT)
    if catatan:
        s.notes_slide.notes_text_frame.text = catatan
    return s


EKA, ERLIS, SRI = "Eka Putra Kurnia", "Erlisnawati", "Sri Rejeki"

# =============== 1. Judul ===============
s = slide_baru(judul_slide=True, catatan=(
    "[Eka · ±0:15] Assalamualaikum, selamat pagi/siang. Kami kelompok Eka Putra Kurnia, Erlisnawati, dan Sri Rejeki. "
    "Kami menyusun rencana proyek Sistem Peringatan Dini Cuaca Publik, disingkat SPDC, yang terdiri atas tiga dokumen "
    "yang saling terhubung: WBS, jadwal, dan Risk Register, mengikuti PMBOK."))
teks(s, 0.9, 1.3, 10.4, 0.5, "TUGAS KELOMPOK — IT PROJECT MANAGEMENT", size=16, bold=True, color=BIRU)
teks(s, 0.9, 1.85, 10.4, 1.6, ["Rencana Proyek Sistem Peringatan Dini", "Cuaca Publik (SPDC)"],
     size=34, bold=True, color=NAVY, font=JUDUL_FONT, space_after=0)
teks(s, 0.9, 3.5, 11.5, 0.8,
     "Sistem IoT audio-visual yang menyalakan panel LED, lampu rotari, dan sirene saat BMKG menerbitkan peringatan dini",
     size=18, italic=True, color=ABU)
for i, (nama, arti) in enumerate([("WBS", "apa saja pekerjaannya"), ("Jadwal", "kapan & berapa lama"),
                                  ("Risk Register", "apa yang bisa salah")]):
    x = 0.9 + i * 2.55
    kotak(s, x, 4.65, 2.35, 1.05, fill=BIRU_MUDA)
    teks(s, x + 0.1, 4.72, 2.15, 0.95, [[(nama, {"bold": True, "color": NAVY, "size": 18})],
                                        [(arti, {"italic": True, "color": BIRU})]],
         size=14, align=PP_ALIGN.CENTER, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
teks(s, 9.0, 4.45, 3.9, 2.6, [
    [("Disusun oleh:", {"bold": True})],
    "EKA PUTRA KURNIA — 25.55.2927",
    "ERLISNAWATI — 25.55.2882",
    "SRI REJEKI — 25.55.2906",
    [("S2 Informatika — Universitas Amikom Yogyakarta", {"color": ABU})],
    [("Dosen: Melwin Syafrizal, S.Kom., M.Eng., Ph.D.", {"color": ABU})],
], size=13, color=TINTA, space_after=2)

# =============== 2. Proyek SPDC ===============
s = slide_baru("SPDC: dari halaman BMKG ke sirene di lapangan", penyaji=EKA, catatan=(
    "[Eka · ±1:00] BMKG menerbitkan peringatan dini cuaca per provinsi di situsnya, tetapi warga di pasar, terminal, atau "
    "bantaran sungai sering tidak membacanya tepat waktu. SPDC menjembatani itu. Server membuka halaman BMKG setiap "
    "5 menit dengan Playwright dan mencari nama wilayah sasaran. Bila ketemu, server mengirim perintah lewat MQTT hanya ke "
    "perangkat di wilayah itu. Perangkat ESP32 menyalakan panel LED, lampu rotari, dan sirene. Peringatan selesai bila nama "
    "wilayah tidak muncul lagi pada 3 pengecekan berturut-turut. Admin memantau semua perangkat dari dasbor dan bisa "
    "menyalakan peringatan secara manual. Targetnya: 6 perangkat di 3 wilayah DIY, menyala paling lambat 10 menit, "
    "selesai dalam sekitar 4 bulan. Sponsornya BPBD."))
alur = [
    ("Halaman BMKG", "Peringatan dini cuaca per provinsi"),
    ("Server", "Playwright membuka halaman tiap 5 menit, mencari nama wilayah"),
    ("MQTT", "Perintah hanya ke perangkat di wilayah itu"),
    ("ESP32", "Perangkat IoT di tiap lokasi (firmware C++)"),
    ("LED · Lampu rotari · Sirene", "Peringatan audio-visual untuk warga"),
]
for i, (h, d) in enumerate(alur):
    x = 0.6 + i * 2.5
    akhir = i == len(alur) - 1
    kotak(s, x, 1.5, 2.15, 1.9, fill=MERAH_MUDA if akhir else BIRU_MUDA)
    teks(s, x + 0.1, 1.6, 1.95, 0.75, h, size=17, bold=True, color=MERAH if akhir else NAVY,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    teks(s, x + 0.1, 2.35, 1.95, 1.0, d, size=13, align=PP_ALIGN.CENTER)
    if not akhir:
        panah(s, x + 2.18, 2.3, w=0.29, h=0.3)
kotak(s, 3.1, 3.75, 7.15, 0.95, fill=PUTIH, line=BIRU)
teks(s, 3.25, 3.8, 6.85, 0.85, [
    [("Dasbor admin  ", {"bold": True, "color": BIRU}),
     ("◄ heartbeat & diagnosa dari setiap perangkat", {})],
    "Admin bisa menyalakan peringatan secara manual bila pembacaan otomatis terganggu",
], size=13, align=PP_ALIGN.CENTER, space_after=2, anchor=MSO_ANCHOR.MIDDLE)
info = [
    ("Tujuan (±4 bulan)", "6 perangkat di 3 wilayah DIY menyala ≤ 10 menit setelah wilayahnya muncul di BMKG"),
    ("Sponsor & tim", "BPBD sebagai pendana dan operator; tim proyek 5 orang (skenario)"),
    ("Di luar lingkup", "Prakiraan cuaca sendiri, sensor cuaca, aplikasi HP, perawatan setelah serah terima"),
]
for i, (h, d) in enumerate(info):
    x = 0.6 + i * 4.1
    kotak(s, x, 5.0, 3.9, 1.75, fill=ABU_MUDA)
    teks(s, x + 0.2, 5.1, 3.5, 1.6, [[(h, {"bold": True, "color": NAVY, "size": 16})], d], size=14, space_after=4)

# =============== 3. Metodologi ===============
s = slide_baru("Kerangka PMBOK: tiga dokumen, satu rencana", penyaji=EKA, catatan=(
    "[Eka · ±0:30] Kami memakai proses perencanaan PMBOK. Dari knowledge area Scope kami membuat WBS; dari Time, jadwal "
    "dengan estimasi PERT dan Critical Path Method; dari Risk, Risk Register dengan matriks kemungkinan kali dampak. Ketiganya "
    "dihubungkan oleh kode WBS yang sama. Pendekatannya prediktif atau Waterfall, karena pembelian, perakitan, dan pemasangan "
    "perangkat keras harus berurutan. WBS dan jadwal juga kami susun di ClickUp."))
kolom = [
    ("Scope", "WBS", ["Create WBS: 3 level", "Aturan 100%", "Kamus WBS + penanggung jawab"]),
    ("Time / Schedule", "Jadwal & Gantt", ["Urutan: hubungan Finish-to-Start", "Durasi: estimasi PERT (3 titik)",
                                           "Jalur kritis: Critical Path Method"]),
    ("Risk", "Risk Register", ["10 risiko, kategori RBS", "Matriks kemungkinan × dampak 5×5",
                               "5 strategi respons + pemicu"]),
]
for i, (ka, dok, items) in enumerate(kolom):
    x = 0.6 + i * 4.1
    kotak(s, x, 1.45, 3.9, 3.55, fill=BIRU_MUDA)
    teks(s, x + 0.25, 1.6, 3.4, 0.4, f"Knowledge area: {ka}", size=13, color=BIRU, bold=True)
    teks(s, x + 0.25, 2.0, 3.4, 0.6, dok, size=24, bold=True, color=NAVY, font=JUDUL_FONT)
    teks(s, x + 0.25, 2.75, 3.4, 2.2, ["• " + t for t in items], size=15, space_after=6)
    if i < 2:
        panah(s, x + 3.92, 3.05, w=0.16, h=0.3, arah=MSO_SHAPE.RIGHT_ARROW, fill=BIRU)
kotak(s, 0.6, 5.25, 12.1, 1.45, fill=ABU_MUDA)
teks(s, 0.85, 5.35, 11.6, 1.3, [
    [("Dihubungkan oleh kode WBS yang sama: ", {"bold": True, "color": NAVY}),
     ("risiko dicari dari work package dan jalur kritis; respons risiko dikembalikan ke WBS dan jadwal.", {})],
    [("Pendekatan prediktif (Waterfall): ", {"bold": True, "color": NAVY}),
     ("pembelian, perakitan, dan pemasangan perangkat keras harus berurutan.  ", {}),
     ("Alat digital: ", {"bold": True, "color": NAVY}), ("ClickUp (WBS + Gantt + jalur kritis).", {})],
], size=15, space_after=6)

# =============== 4. WBS ===============
s = slide_baru(f"WBS: 3 level, {len(FASE)} fase, {JUMLAH_WP} work package", penyaji=EKA, catatan=(
    f"[Eka · ±1:30] Ini WBS proyek. Level 1 adalah proyek SPDC, level 2 adalah {len(FASE)} fase, dan level 3 adalah "
    f"{JUMLAH_WP} work package, yaitu pekerjaan terkecil yang bisa ditugaskan, diperkirakan durasinya, dan dipantau. WBS "
    "disusun per fase seperti contoh PMBOK edisi 4 Gambar 5-9. Kami memegang aturan 100 persen: semua pekerjaan yang "
    "diperlukan harus masuk, termasuk manajemen proyek, pengadaan dan izin, sampai pelatihan operator. Setiap work package "
    "punya uraian, hasil, dan penanggung jawab di kamus WBS, Lampiran A laporan. WBS menunjukkan pemecahan pekerjaan, "
    "bukan urutan; urutannya ada di jadwal."))
puncak = kotak(s, 3.9, 1.25, 5.5, 0.55, fill=NAVY, radius=0.15)
label(puncak, "1  Sistem Peringatan Dini Cuaca Publik (SPDC)", size=15, color=PUTIH, bold=True)
lebar_f, jarak_f, x0 = 1.48, 0.07, 0.6
garis(s, 6.64, 1.8, 0.025, 0.18)
garis(s, x0 + lebar_f / 2, 1.97, (len(FASE) - 1) * (lebar_f + jarak_f), 0.025)
for j, (kode, nama, wps) in enumerate(FASE):
    x = x0 + j * (lebar_f + jarak_f)
    garis(s, x + lebar_f / 2 - 0.0125, 1.97, 0.025, 0.18)
    fase = kotak(s, x, 2.15, lebar_f, 0.95, fill=BIRU, radius=0.1)
    label(fase, [kode, nama], size=11, color=PUTIH, bold=True)
    kotak(s, x, 3.17, lebar_f, 3.4, fill=BIRU_MUDA, radius=0.04)
    teks(s, x + 0.03, 3.22, lebar_f - 0.06, 3.3,
         [[(k + " ", {"bold": True, "color": NAVY}), (n, {})] for k, n in wps], size=10, space_after=3)
teks(s, 0.6, 6.6, 12.1, 0.35, [[
    ("Level 1 ", {"bold": True, "color": NAVY}), ("proyek  ·  ", {}),
    ("Level 2 ", {"bold": True, "color": BIRU}), ("fase  ·  ", {}),
    ("Level 3 ", {"bold": True, "color": NAVY}), ("work package (paket kerja)  ·  ", {}),
    ("Aturan 100%: ", {"bold": True}), ("semua pekerjaan masuk, tidak ada yang di luar WBS", {})]],
    size=12, color=ABU, align=PP_ALIGN.CENTER)

# =============== 5. Pekerjaan yang sering terlupa ===============
s = slide_baru("Pekerjaan yang sering terlupa juga masuk WBS", penyaji=EKA, catatan=(
    "[Eka · ±0:30] Contoh penerapan aturan 100 persen. Selain membuat server dan perangkat, WBS memuat manajemen proyek, "
    "izin pemasangan, deploy dan pengamanan server, sosialisasi warga, dan pelatihan operator. Perbaikan bug tidak dibuat "
    "work package terpisah; ia masuk ke fase pengujian dengan pola uji, perbaiki, uji ulang. Milestone tidak masuk WBS "
    "karena bukan pekerjaan, melainkan penanda waktu di jadwal. Selanjutnya, jadwal akan dijelaskan oleh Erlisnawati."))
kartu = [
    ("1.1", "Manajemen proyek", "Project charter, rapat mingguan, pengendalian perubahan, penutupan"),
    ("1.4.4", "Izin pemasangan", "Dari pemilik lokasi dan pemda; paling tidak pasti (10–30 hari)"),
    ("1.8.2", "Deploy & pengamanan", "Server dipasang online dengan firewall dan HTTPS"),
    ("1.8.4", "Sosialisasi warga", "Arti setiap bunyi dan lampu, serta tindakan warga"),
    ("1.8.5", "Pelatihan operator", "Petugas BPBD memakai dasbor dan mode uji"),
    ("1.7.x", "Perbaikan bug", "Masuk 31 hari pengujian: uji → perbaiki → uji ulang"),
]
for i, (kode, h, d) in enumerate(kartu):
    x = 0.6 + (i % 3) * 4.1
    y = 1.45 + (i // 3) * 2.35
    kotak(s, x, y, 3.9, 2.15, fill=BIRU_MUDA)
    teks(s, x + 0.25, y + 0.15, 3.4, 0.45, kode, size=15, bold=True, color=BIRU)
    teks(s, x + 0.25, y + 0.55, 3.4, 0.5, h, size=19, bold=True, color=NAVY)
    teks(s, x + 0.25, y + 1.1, 3.4, 1.0, d, size=14)
teks(s, 0.6, 6.25, 12.1, 0.5, [[("Milestone tidak masuk WBS: ", {"bold": True}),
                                 ("milestone bukan pekerjaan, melainkan penanda waktu di jadwal.", {})]],
     size=14, color=ABU)

# =============== 6. Estimasi dan ringkasan jadwal ===============
contoh = AKTIVITAS["1.4.2"]
s = slide_baru("Estimasi durasi dan ringkasan jadwal", penyaji=ERLIS, catatan=(
    "[Erlisnawati · ±1:00] Terima kasih. Setiap work package menjadi satu aktivitas. Durasinya kami perkirakan dengan "
    "estimasi tiga titik atau PERT: optimis ditambah empat kali paling mungkin ditambah pesimis, dibagi enam. Contohnya "
    "pembelian komponen: 7, 10, dan 20 hari menghasilkan 11 hari. Urutannya disusun dengan hubungan Finish-to-Start, total "
    f"65 ketergantungan. Hasilnya: proyek {TOTAL} hari kerja atau sekitar 4,2 bulan, dari 4 Januari sampai 11 Mei 2027, "
    "dengan 6 milestone yang masing-masing disetujui sponsor atau manajer proyek."))
kotak(s, 0.6, 1.45, 5.7, 2.95, fill=BIRU_MUDA)
teks(s, 0.85, 1.55, 5.2, 2.8, [
    [("Estimasi tiga titik (PERT)", {"bold": True, "color": NAVY, "size": 18})],
    [("Durasi = (O + 4M + P) / 6", {"bold": True, "color": BIRU, "size": 22, "font": JUDUL_FONT})],
    "O = optimis · M = paling mungkin · P = pesimis",
    [("Contoh 1.4.2 Pembelian komponen:", {"bold": True})],
    "(7 + 4×10 + 20) / 6 = 11,17 → 11 hari kerja",
    [("Rentangnya lebar (7–20 hari) → jadi risiko R-01", {"italic": True, "color": MERAH})],
], size=15, space_after=5)
stat = [
    (f"{TOTAL} hari kerja", "≈ 18,4 minggu ≈ 4,2 bulan"),
    (f"{tanggal(1)} – {tanggal(TOTAL)}", "2027; Senin–Jumat"),
    (f"{JUMLAH_WP} + {len(MILESTONE)}", "work package + milestone"),
    ("65", "ketergantungan Finish-to-Start"),
]
for i, (big, small) in enumerate(stat):
    x = 6.55 + (i % 2) * 3.1
    y = 1.45 + (i // 2) * 1.5
    kotak(s, x, y, 2.95, 1.4, fill=ABU_MUDA)
    teks(s, x + 0.15, y + 0.12, 2.65, 0.65, big, size=22, bold=True, color=NAVY, font=JUDUL_FONT)
    teks(s, x + 0.15, y + 0.8, 2.65, 0.5, small, size=13, color=ABU)
# garis waktu milestone
gx, gw, gy = 0.9, 11.5, 5.4
garis(s, gx, gy, gw, 0.05, fill=NAVY)
teks(s, 0.6, 4.55, 4.0, 0.4, "Milestone (hari kerja ke-)", size=13, bold=True, color=NAVY)
nama_m = {"M1": "Proyek disahkan", "M2": "Kebutuhan disetujui", "M3": "Desain disetujui",
          "M4": "UAT disetujui", "M5": "Serah terima", "M6": "Proyek ditutup"}
for k in MILESTONE:
    hari = AKTIVITAS[k]["ef"]
    x = gx + hari / TOTAL * gw
    kotak(s, x - 0.11, gy - 0.085, 0.22, 0.22, fill=MERAH if AKTIVITAS[k]["fl"] == 0 else BIRU, bentuk=MSO_SHAPE.DIAMOND)
    teks(s, x - 0.3, gy - 0.45, 0.6, 0.35, k, size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
teks(s, 0.6, 5.75, 12.1, 0.9, [
    "   ".join(f"{k} {nama_m[k]} ({AKTIVITAS[k]['ef']}; {tanggal(AKTIVITAS[k]['ef'])})" for k in MILESTONE[:3]),
    "   ".join(f"{k} {nama_m[k]} ({AKTIVITAS[k]['ef']}; {tanggal(AKTIVITAS[k]['ef'])})" for k in MILESTONE[3:]),
], size=13, color=TINTA, align=PP_ALIGN.CENTER, space_after=2)

# =============== 7. Gantt dan jalur kritis ===============
kritis = [k for k, a in AKTIVITAS.items() if a["fl"] == 0 and a["dur"] > 0]
s = slide_baru("Gantt Chart dan jalur kritis (ClickUp)", penyaji=ERLIS, catatan=(
    "[Erlisnawati · ±1:30] Ini Gantt Chart di ClickUp, bagian Februari sampai Mei. Batang merah adalah jalur kritis, "
    "yaitu rangkaian aktivitas dengan float nol: bila satu saja terlambat, seluruh proyek mundur. Jalur kritisnya melewati "
    "pembelian komponen, prototipe dan firmware, uji perangkat keras, lalu bercabang dua dengan panjang sama: cabang perakitan "
    "dan pemasangan, serta cabang uji integrasi dan uji kegagalan. Keduanya bertemu di UAT, lalu serah terima dan penutupan. "
    "Perangkat lunak server tidak kritis; masih punya kelonggaran 10 hari. ClickUp menghitung jalur kritis sendiri, dan "
    f"hasilnya sama dengan perhitungan CPM kami: {len(kritis)} work package."))
pic = s.shapes.add_picture(str(GANTT), Inches(0.5), Inches(1.3), height=Inches(5.5))
# potong: buang menu ClickUp di kiri dan atas serta bulan Januari yang kosong; sisakan sumbu waktu Feb–Mei
pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = 690 / 1920, 20 / 1920, 108 / 919, 30 / 919
pic.width = int(pic.height * (1920 * (1 - pic.crop_left - pic.crop_right)) / (919 * (1 - pic.crop_top - pic.crop_bottom)))
pic.line.color.rgb = REDUP
kanan = 0.5 + pic.width / 914400 + 0.25
lebar_k = 12.75 - kanan
kotak(s, kanan, 1.3, lebar_k, 5.5, fill=ABU_MUDA)
langkah = [
    ("1.4.1–1.4.2", "Daftar & beli komponen"),
    ("1.6.1–1.6.4", "Prototipe dan firmware"),
    ("1.7.3", "Uji perangkat keras"),
    ("A: 1.6.5 · 1.8.1 · 1.8.3", "Rakit, pasang, konfigurasi"),
    ("B: 1.7.4 · 1.7.5", "Uji integrasi & uji kegagalan"),
    ("1.7.6 → 1.8.6 → 1.1.4", "UAT, serah terima, penutupan"),
]
teks(s, kanan + 0.15, 1.4, lebar_k - 0.3, 0.45, "Rantai merah = jalur kritis", size=16, bold=True, color=MERAH)
for i, (k, d) in enumerate(langkah):
    y = 1.9 + i * 0.62
    teks(s, kanan + 0.15, y, lebar_k - 0.3, 0.6, [[(k, {"bold": True, "color": NAVY, "size": 12})], d],
         size=12, space_after=0)
teks(s, kanan + 0.15, 5.65, lebar_k - 0.3, 1.1, [
    [("Float 0: ", {"bold": True}), ("terlambat sehari = proyek mundur sehari", {})],
    [("Server: ", {"bold": True}), ("float 10 hari, tidak kritis", {})],
], size=12, space_after=2)

# =============== 8. Cadangan dan pemendekan ===============
s = slide_baru("Cadangan waktu dan pilihan mempercepat", penyaji=ERLIS, catatan=(
    "[Erlisnawati · ±0:30] Manajer proyek memegang cadangan waktu 5 hari kerja, terutama untuk bug yang lebih banyak "
    "dari perkiraan. Bila terpakai semua, proyek selesai paling lambat hari ke-97. Bila sponsor ingin lebih cepat, ada "
    "pilihan fast tracking: komponen prototipe dibeli lebih dulu. Hasilnya 88 hari dengan satu programmer atau 82 hari dengan "
    "dua, tetapi tidak kami terapkan karena 92 hari sudah sesuai target. Kami juga mengecek beban per peran: dengan satu "
    "programmer server pun hasilnya tetap 92 hari. Berikutnya Sri Rejeki akan membahas risiko."))
kartu = [
    ("Cadangan waktu", "5 hari kerja", BIRU, BIRU_MUDA,
     ["Contingency reserve, dipegang manajer proyek", f"Tidak masuk jadwal; batas terburuk hari ke-{TOTAL + 5}",
      "Untuk bug yang lebih banyak dari perkiraan (R-06)"]),
    ("Fast tracking", "88 / 82 hari", NAVY, ABU_MUDA,
     ["Komponen prototipe dibeli lebih dulu", "88 hari (1 programmer) atau 82 hari (2 programmer)",
      f"Tidak diterapkan: {TOTAL} hari sudah sesuai target"]),
    ("Cek sumber daya", f"Tetap {TOTAL} hari", NAVY, ABU_MUDA,
     ["Tim 5 orang; Engineer IoT merangkap teknisi lapangan", "Pekerjaan yang bertumpuk masih muat dalam float",
      "Satu programmer server pun cukup"]),
]
for i, (h, big, warna, latar, items) in enumerate(kartu):
    x = 0.6 + i * 4.1
    kotak(s, x, 1.45, 3.9, 4.6, fill=latar)
    teks(s, x + 0.25, 1.65, 3.4, 0.45, h, size=17, bold=True, color=warna)
    teks(s, x + 0.25, 2.15, 3.4, 0.8, big, size=30, bold=True, color=warna, font=JUDUL_FONT)
    teks(s, x + 0.25, 3.15, 3.4, 3.4, ["• " + t for t in items], size=15, space_after=10)

# =============== 9. Risk Register dan matriks ===============
awal = {lv: sum(1 for r in RISIKO if level(r["k"] * r["d"]) == lv) for lv in WARNA_LEVEL}
sisa = {lv: sum(1 for r in RISIKO if level(r["k2"] * r["d2"]) == lv) for lv in WARNA_LEVEL}
s = slide_baru(f"Risk Register: {len(RISIKO)} risiko dalam matriks 5 × 5", penyaji=SRI, catatan=(
    "[Sri Rejeki · ±1:15] Terima kasih. Kami mengidentifikasi 10 risiko dengan menelusuri WBS, jalur kritis, dan kategori "
    "risiko atau RBS. Setiap risiko dinilai kemungkinan dan dampaknya dengan skala 1 sampai 5; skala kemungkinan mengikuti "
    "contoh Risk Register dari Bapak, dan skala dampak mengikuti PMBOK. Nilai risiko adalah kemungkinan kali dampak. "
    f"Sebelum respons ada {awal['Tinggi']} risiko Tinggi dan {awal['Sedang']} Sedang. Setelah respons, tidak ada lagi "
    f"risiko Tinggi: {sisa['Sedang']} Sedang dan {sisa['Rendah']} Rendah. Kelima strategi PMBOK dipakai: avoid, mitigate, "
    "transfer, accept aktif, dan escalate."))


def matriks(x, judul, residual):
    teks(s, x, 1.35, 5.6, 0.4, judul, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    t = s.shapes.add_table(6, 6, Inches(x + 0.35), Inches(1.8), Inches(5.25), Inches(3.6)).table
    t.columns[0].width = Inches(0.5)
    for j in range(1, 6):
        t.columns[j].width = Inches(0.95)
    for i in range(6):
        t.rows[i].height = Inches(0.6)
    tabel_sel(t.cell(0, 0), "K \\ D", size=10, bold=True, fill=ABU_MUDA, align=PP_ALIGN.CENTER)
    for d in range(1, 6):
        tabel_sel(t.cell(0, d), str(d), size=12, bold=True, fill=ABU_MUDA, align=PP_ALIGN.CENTER)
    for b, k in enumerate(range(5, 0, -1), start=1):
        tabel_sel(t.cell(b, 0), str(k), size=12, bold=True, fill=ABU_MUDA, align=PP_ALIGN.CENTER)
        for d in range(1, 6):
            ids = [r["id"] for r in RISIKO if ((r["k2"], r["d2"]) if residual else (r["k"], r["d"])) == (k, d)]
            isi = [[(str(k * d), {"size": 9, "color": ABU})]]
            if ids:
                isi.append([(", ".join(i_.replace("R-", "") for i_ in ids),
                             {"bold": True, "size": 12 if len(ids) <= 3 else 10})])
            tabel_sel(t.cell(b, d), isi, fill=WARNA_LEVEL[level(k * d)], align=PP_ALIGN.CENTER)
    teks(s, x, 3.2, 0.35, 0.6, "K", size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


matriks(0.6, "Sebelum respons", False)
matriks(7.15, "Sesudah respons (residual)", True)
panah(s, 6.3, 3.4, w=0.55, h=0.4, fill=BIRU)
teks(s, 0.6, 5.45, 12.1, 0.35, "K = kemungkinan (1–5)  ·  D = dampak (1–5)  ·  angka tebal = nomor risiko (R-xx)  ·  "
     "Rendah 1–5 · Sedang 6–12 · Tinggi 13–19 · Kritikal 20–25", size=12, color=ABU, align=PP_ALIGN.CENTER)
kotak(s, 0.6, 5.85, 12.1, 1.0, fill=ABU_MUDA)
teks(s, 0.85, 5.9, 11.6, 0.9, [
    [("Sebelum: ", {"bold": True}), (f"{awal['Tinggi']} Tinggi, {awal['Sedang']} Sedang  →  ", {}),
     ("Sesudah: ", {"bold": True}), (f"0 Tinggi, {sisa['Sedang']} Sedang, {sisa['Rendah']} Rendah", {})],
    [("Strategi: ", {"bold": True}), ("Avoid (hindari) · Mitigate (kurangi) · Transfer (alihkan) · "
                                      "Accept aktif (terima + rencana darurat) · Escalate (serahkan ke yang berwenang)", {})],
], size=14, space_after=2, anchor=MSO_ANCHOR.MIDDLE)

# =============== 10. Tiga risiko teratas ===============
s = slide_baru("Tiga risiko utama dan responsnya", penyaji=SRI, catatan=(
    "[Sri Rejeki · ±1:30] Tiga risiko yang kami soroti. Pertama, R-01: komponen terlambat datang. Pembelian komponen ada "
    "di jalur kritis, jadi keterlambatan 10 hari langsung memundurkan proyek. Responsnya menghindari barang impor, menyiapkan "
    "dua pemasok, dan membeli unit cadangan. Kedua, R-03: situs BMKG berubah atau memblokir akses otomatis, sehingga fungsi "
    "utama berhenti. Responsnya pembacaan yang tidak bergantung pada bentuk halaman, pengecekan yang tidak membebani server "
    "BMKG, dan rencana darurat berupa aktivasi manual dari dasbor. Ketiga, R-05: listrik atau internet putus saat cuaca "
    "buruk, justru saat sirene paling dibutuhkan. Ini tidak bisa diatasi dengan update perangkat lunak, sehingga responsnya "
    "baterai cadangan minimal 4 jam, modem 4G, dan simulasi padam saat uji kegagalan. Setiap risiko punya pemicu yang jelas "
    "dan pemiliknya."))
tiga = [
    ("R-01", "Komponen terlambat datang", "K4 × D4 = 16 Tinggi", "→ 6 Sedang",
     "Stok kosong atau barang impor → telat ±10 hari → 1.4.2 di jalur kritis, proyek mundur",
     "Avoid + Mitigate", ["Utamakan pemasok dalam negeri (ESP32 mudah didapat)", "Minimal 2 pemasok per komponen",
                          "Pesan segera + 1 unit cadangan per jenis"],
     "Stok belum dikonfirmasi 2 hari kerja setelah pesan", "Manajer Proyek"),
    ("R-03", "Situs BMKG berubah / memblokir", "K3 × D5 = 15 Tinggi", "→ 6 Sedang",
     "Tampilan diubah atau akses otomatis dibatasi → peringatan tak terbaca → fungsi utama berhenti",
     "Mitigate + Accept aktif", ["Cari teks nama wilayah, tak bergantung bentuk halaman",
                                 "Cek tiap 5 menit, hanya provinsi yang relevan", "Darurat: aktivasi manual dari dasbor"],
     "Gagal muat 3 kali berturut-turut", "Programmer Server"),
    ("R-05", "Listrik / internet putus", "K3 × D4 = 12 Sedang", "→ 6 Sedang",
     "Padam saat cuaca buruk → perangkat tak menerima perintah, justru saat paling dibutuhkan",
     "Mitigate", ["Baterai cadangan ≥ 4 jam", "Modem 4G bila Wi-Fi mati; sambung ulang otomatis",
                  "Simulasi padam di uji kegagalan (1.7.5)"],
     "Baterai < 4 jam atau heartbeat hilang > 5 menit saat uji", "Engineer IoT"),
]
for i, (rid, h, nilai, sisa_r, apa, strat, items, pemicu, pemilik) in enumerate(tiga):
    x = 0.6 + i * 4.1
    tinggi = nilai.endswith("Tinggi")
    kotak(s, x, 1.4, 3.9, 5.4, fill=MERAH_MUDA if tinggi else BIRU_MUDA)
    teks(s, x + 0.2, 1.5, 3.5, 0.75, [[(rid + "  ", {"bold": True, "color": MERAH if tinggi else BIRU}),
                                       (h, {"bold": True, "color": NAVY})]], size=17)
    teks(s, x + 0.2, 2.2, 3.5, 0.4, [[(nilai + " ", {"bold": True}), (sisa_r, {"color": ABU})]], size=13)
    teks(s, x + 0.2, 2.6, 3.5, 1.0, apa, size=12, italic=True, color=ABU)
    teks(s, x + 0.2, 3.6, 3.5, 0.4, strat, size=14, bold=True, color=NAVY)
    teks(s, x + 0.2, 4.0, 3.5, 1.75, ["• " + t for t in items], size=13, space_after=4)
    teks(s, x + 0.2, 5.75, 3.5, 1.0, [[("Pemicu: ", {"bold": True}), (pemicu, {})],
                                      [("Pemilik: ", {"bold": True}), (pemilik, {})]], size=11.5, space_after=2)

# =============== 11. Keterkaitan dan kesimpulan ===============
s = slide_baru("Keterkaitan: respons risiko kembali ke WBS dan jadwal", penyaji=SRI, catatan=(
    "[Sri Rejeki · ±0:30] Ketiga dokumen terhubung lewat kode WBS. Contohnya, R-01 menyangkut 1.4.2 yang kritis; responsnya "
    "masuk ke pekerjaan BOM dan pembelian. R-05 menambah baterai dan modem 4G ke daftar komponen. R-06 menghasilkan cadangan "
    "waktu 5 hari. R-03 menambah aktivasi manual ke dasbor. Kesimpulannya, lama proyek ditentukan oleh rantai perangkat keras "
    "dan pengujian akhir, dua risiko teratas berada di rantai itu, dan setelah respons tidak ada lagi risiko Tinggi. "
    "Terima kasih; kami siap menjawab pertanyaan."))
t = s.shapes.add_table(5, 3, Inches(0.6), Inches(1.4), Inches(12.1), Inches(3.4)).table
for j, w in enumerate((3.2, 3.4, 5.5)):
    t.columns[j].width = Inches(w)
for j, h in enumerate(("Risiko", "WBS dan posisi di jadwal", "Respons yang masuk kembali ke rencana")):
    tabel_sel(t.cell(0, j), h, size=14, bold=True, color=PUTIH, fill=NAVY)
baris_kait = [
    ("R-01 Komponen terlambat", "1.4.2 (kritis)", "Pemasok dalam negeri, 2 pemasok, unit cadangan di 1.4.1–1.4.2; opsi fast tracking"),
    ("R-03 Situs BMKG berubah", "1.2.2, 1.5.1 (float 5–10)", "Aktivasi manual ditambahkan ke dasbor (1.3.6, 1.5.5)"),
    ("R-05 Listrik/internet putus", "1.3.4, 1.6.3, 1.7.5 (kritis)", "Baterai dan modem 4G masuk daftar komponen (1.4.1); simulasi padam di 1.7.5"),
    ("R-06 Bug lebih banyak", "1.7.4, 1.7.5 (kritis)", f"Uji dini 1.7.1–1.7.2; cadangan waktu 5 hari (batas hari ke-{TOTAL + 5})"),
]
for i, b in enumerate(baris_kait, start=1):
    for j, v in enumerate(b):
        tabel_sel(t.cell(i, j), v, size=13, bold=(j == 0), color=NAVY if j == 0 else TINTA,
                  fill=BIRU_MUDA if i % 2 else PUTIH)
kotak(s, 0.6, 5.0, 12.1, 1.8, fill=ABU_MUDA)
teks(s, 0.85, 5.1, 8.6, 1.65, [
    [("Kesimpulan", {"bold": True, "color": NAVY, "size": 17})],
    f"• {TOTAL} hari kerja; lama proyek ditentukan rantai perangkat keras dan pengujian akhir",
    "• Dua risiko teratas (R-01, R-03) berada di rantai itu atau di fungsi utama",
    "• Setelah respons, tidak ada risiko Tinggi; cadangan 5 hari untuk bug",
], size=14, space_after=3)
teks(s, 9.6, 5.1, 3.0, 1.65, "Terima kasih", size=30, bold=True, color=BIRU, font=JUDUL_FONT,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# =============== C1. Jalur kritis (cadangan) ===============
s = slide_baru("Cadangan: aktivitas di jalur kritis (CPM)", catatan=(
    "Slide cadangan untuk menjawab pertanyaan; tidak dipresentasikan. ES/EF = mulai/selesai paling awal (hari kerja ke-). "
    "Float = LS − ES; nol berarti kritis. Tabel lengkap ada di Lampiran B laporan."))
urut = sorted(kritis, key=lambda k: (AKTIVITAS[k]["es"], AKTIVITAS[k]["ef"]))
for kol, bagian in enumerate((urut[:11], urut[11:])):
    t = s.shapes.add_table(len(bagian) + 1, 4, Inches(0.6 + kol * 6.2), Inches(1.35), Inches(5.9),
                           Inches(0.4 * (len(bagian) + 1))).table
    for j, w in enumerate((0.7, 3.65, 0.7, 0.85)):
        t.columns[j].width = Inches(w)
    for j, h in enumerate(("Kode", "Aktivitas", "Durasi", "ES–EF")):
        tabel_sel(t.cell(0, j), h, size=12, bold=True, color=PUTIH, fill=NAVY)
    for i, k in enumerate(bagian, start=1):
        a = AKTIVITAS[k]
        for j, v in enumerate((k, a["nama"], str(a["dur"]), f"{a['es']}–{a['ef']}")):
            tabel_sel(t.cell(i, j), v, size=11, bold=(j == 0), fill=MERAH_MUDA if i % 2 else PUTIH)
teks(s, 0.6, 6.45, 12.1, 0.4, f"{len(kritis)} work package dengan float 0 + milestone M1, M2, M4, M5, M6. "
     "Setelah 1.7.3 jalur bercabang: A (1.6.5 → 1.8.1 → 1.8.3) dan B (1.7.4 → 1.7.5), sama-sama 12 hari.",
     size=12, color=ABU)

# =============== C2. Daftar istilah (cadangan) ===============
s = slide_baru("Cadangan: daftar istilah", catatan="Slide cadangan untuk menjawab pertanyaan; tidak dipresentasikan.")
glos = [
    ("WBS", "pemecahan seluruh pekerjaan proyek secara bertingkat"),
    ("Work package", "pekerjaan terkecil di WBS: bisa ditugaskan & dipantau"),
    ("Aturan 100%", "semua pekerjaan proyek harus masuk WBS"),
    ("Milestone", "titik pencapaian penting, durasi nol"),
    ("PERT", "estimasi durasi (O + 4M + P) / 6"),
    ("Finish-to-Start", "aktivitas mulai setelah pendahulunya selesai"),
    ("CPM / jalur kritis", "rangkaian aktivitas terpanjang, float nol"),
    ("Float", "berapa hari boleh terlambat tanpa menunda proyek"),
    ("Contingency reserve", "cadangan waktu untuk menghadapi risiko"),
    ("Fast tracking", "mempercepat dengan mengerjakan paralel"),
    ("Risk Register", "daftar risiko, nilai, dan responsnya"),
    ("Residual", "nilai risiko sisa setelah respons"),
    ("Avoid / Mitigate", "hindari ancaman / kurangi kemungkinan & dampak"),
    ("Transfer / Escalate", "alihkan ke pihak ketiga / serahkan ke yang berwenang"),
    ("Accept aktif", "terima, tetapi siapkan cadangan atau rencana darurat"),
    ("Playwright", "pustaka Python pengendali peramban web otomatis"),
    ("MQTT", "protokol kirim pesan ringan untuk perangkat IoT"),
    ("ESP32", "papan mikrokontroler dengan Wi-Fi, diprogram C++"),
    ("Heartbeat", "sinyal berkala bahwa perangkat masih hidup"),
    ("UAT", "uji penerimaan oleh sponsor"),
    ("BOM", "daftar komponen beserta jumlah dan spesifikasinya"),
    ("BPBD", "Badan Penanggulangan Bencana Daerah (sponsor)"),
]
for k, bagian in enumerate((glos[:11], glos[11:])):
    tbl = s.shapes.add_table(len(bagian), 2, Inches(0.6 + k * 6.2), Inches(1.4), Inches(5.95), Inches(5.2)).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(2.1), Inches(3.85)
    for i, (t_, d) in enumerate(bagian):
        tabel_sel(tbl.cell(i, 0), t_, size=12, bold=True, color=NAVY, fill=BIRU_MUDA if i % 2 == 0 else PUTIH)
        tabel_sel(tbl.cell(i, 1), d, size=12, fill=BIRU_MUDA if i % 2 == 0 else PUTIH)

# =============== C3. Daftar pustaka ===============
s = slide_baru("Daftar Pustaka", catatan="Slide cadangan; tidak dipresentasikan.")
pustaka = teks_md[teks_md.index("# Daftar Pustaka"):].split("\n#")[0].split("<!--")[0]
teks(s, 0.6, 1.4, 12.1, 5.4, [b.replace("*", "") for b in pustaka.splitlines()[1:] if b.strip()],
     size=13, color=TINTA, space_after=8)

HASIL.mkdir(exist_ok=True)
prs.save(OUT)
print("Tersimpan:", OUT)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HASIL), str(OUT)],
               check=True, capture_output=True)
print("Tersimpan:", OUT.with_suffix(".pdf"))
