"""Membuat laporan Tugas 2 (Rencana Proyek SPDC) dalam format .docx dan .pdf dari laporan_tugas2.md.

Jalankan dengan venv mata kuliah (berisi python-docx):
    ../../.venv/bin/python buat_laporan.py

Hasil: ../hasil/Tugas2_RencanaProyek_SPDC_Laporan.docx dan .pdf (PDF dibuat dengan LibreOffice).

Berkas .md adalah satu-satunya sumber isi. Markup yang dikenali:
    # / ## / ###        judul bab, subbab, sub-subbab
    - teks / 1. teks    butir dan butir bernomor
    | tabel |           tabel Markdown; <!-- tabel: judul | lebar: persen,... --> tepat sebelumnya
    **tebal**, *miring*, ***tebal miring***
    <!-- SAMPUL ... -->                 data sampul (kunci: nilai)
    <!-- LANDSCAPE --> / <!-- PORTRAIT -->   ganti orientasi halaman
    <!-- GANTT -->, <!-- MATRIKS_AWAL -->, <!-- MATRIKS_RESIDUAL -->   dibuat dari tabel jadwal dan Risk Register
    <!-- DIAGRAM_WBS -->                diagram pohon WBS, dibuat dari kamus WBS (Lampiran A)
"""
import math
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

DRAF = Path(__file__).resolve().parent
SUMBER = DRAF / "laporan_tugas2.md"
HASIL = DRAF.parent / "hasil"
OUT = HASIL / "Tugas2_RencanaProyek_SPDC_Laporan.docx"

TOKEN = re.compile(r"(\*\*\*|\*\*|\*)")
WARNA_LEVEL = {"Rendah": "C6E0B4", "Sedang": "FFE699", "Tinggi": "F4B183", "Kritikal": "FF7C80"}
WARNA_KRITIS, WARNA_FLOAT = "C0504D", "9DC3E6"

teks_md = SUMBER.read_text(encoding="utf8")
doc = Document()

# ---------- Pengaturan halaman & gaya ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3
for nama, ukuran in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12), ("Title", 18)):
    st = doc.styles[nama]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(ukuran)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(6)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 3"].font.italic = True


def lebar_isi():
    s = doc.sections[-1]
    return s.page_width - s.left_margin - s.right_margin


def isi_runs(par, teks, ukuran=None):
    """Tulis teks dengan markup ***, **, * (bisa bersarang) ke paragraf."""
    tebal = miring = False
    for bagian in TOKEN.split(teks):
        if bagian == "***":
            tebal, miring = not tebal, not miring
        elif bagian == "**":
            tebal = not tebal
        elif bagian == "*":
            miring = not miring
        elif bagian:
            r = par.add_run(bagian)
            r.bold, r.italic = tebal or None, miring or None
            if ukuran:
                r.font.size = Pt(ukuran)
    return par


def p(teks, align=WD_ALIGN_PARAGRAPH.JUSTIFY, ukuran=None):
    par = doc.add_paragraph()
    par.alignment = align
    return isi_runs(par, teks, ukuran)


def butir(teks, nomor=None):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = par.paragraph_format
    pf.left_indent, pf.first_line_indent, pf.space_after = Cm(0.9), Cm(-0.6), Pt(3)
    isi_runs(par, (f"{nomor}.\t" if nomor else "•\t") + teks)
    pf.tab_stops.add_tab_stop(Cm(0.9))


def judul_tabel(teks):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(6)
    r = par.add_run(teks)
    r.italic = True
    r.font.size = Pt(10)


def arsir(cell, warna):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), warna)
    tcPr.append(shd)


def atur_lebar(t, lebar):
    # Lebar kolom harus ditulis juga di <w:tblGrid> (tblLayout "fixed" sudah diset oleh
    # t.autofit = False), kalau tidak Word Online/LibreOffice membagi kolom sama lebar.
    tblPr = t._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(Emu(x).twips for x in lebar))); tblW.set(qn("w:type"), "dxa")
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), lebar):
        gc.set(qn("w:w"), str(Emu(w).twips))
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = lebar[i]


def ulang_header(t):
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)


def tabel(header, baris, persen=None, judul=None):
    if judul:
        judul_tabel(judul)
    persen = persen or [100 / len(header)] * len(header)
    ukuran = 10 if len(header) <= 5 else 8.5
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    kol_level = [i for i, h in enumerate(header) if h.strip("* ") == "Level"]
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        isi_runs(c.paragraphs[0], f"**{h.replace('**', '')}**", ukuran)
        arsir(c, "E8E8E8")
    for b in baris:
        cells = t.add_row().cells
        for i, v in enumerate(b):
            isi_runs(cells[i].paragraphs[0], v, ukuran)
            if i in kol_level and v.strip("* ") in WARNA_LEVEL:
                arsir(cells[i], WARNA_LEVEL[v.strip("* ")])
    atur_lebar(t, [int(lebar_isi() * x / sum(persen)) for x in persen])
    kecil = len(baris) <= 10    # tabel pendek dijaga tetap di satu halaman
    for n, row in enumerate(t.rows):
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.05
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.keep_with_next = kecil and n < len(t.rows) - 1
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    ulang_header(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def nomor_halaman(section, mulai=None):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
        r._r.append(el)
    r.font.size = Pt(10)
    if mulai:
        pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), str(mulai))
        section._sectPr.append(pg)


def kepala_halaman(section, teks):
    par = section.header.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = par.add_run(teks)
    r.italic = True
    r.font.size = Pt(9)


def section_baru(landscape):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    lebar, tinggi = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = lebar, tinggi
    m = Cm(2) if landscape else Cm(2.5)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, m)
    # add_section menyalin pgNumType dari section isi; tanpa dihapus, nomor halaman mulai lagi dari 1
    for el in s._sectPr.findall(qn("w:pgNumType")):
        s._sectPr.remove(el)


# ---------- Data dari tabel di .md (untuk Gantt dan matriks) ----------
def sel(baris):
    return [c.strip() for c in baris.strip().strip("|").split("|")]


def bersih(c):
    return c.replace("*", "").strip()


def ambil_tabel(kepala):
    """Baris isi (tanpa header dan pemisah) dari tabel pertama yang diawali `kepala`."""
    blok = []
    for b in teks_md[teks_md.index(kepala):].splitlines():
        if not b.startswith("|"):
            break
        blok.append(sel(b))
    return blok[2:]


AKTIVITAS = [dict(kode=bersih(c[0]), nama=bersih(c[1]), dur=int(bersih(c[6])), es=int(c[8]), ef=int(c[9]),
                  fl=int(bersih(c[12]))) for c in ambil_tabel("| Kode | Aktivitas | O |")]
RISIKO = []
for c in ambil_tabel("| ID | Kategori"):
    k2, d2 = map(int, re.match(r"(\d) × (\d)", c[10]).groups())
    RISIKO.append(dict(id=bersih(c[0]), k=int(c[4]), d=int(c[5]), k2=k2, d2=d2))


FASE = []   # [(kode, nama, [(kode_wp, nama_wp), ...]), ...]
for c in ambil_tabel("| Kode | Work package | Uraian |"):
    kode, nama = bersih(c[0]), bersih(c[1])
    if kode.count(".") == 1:
        FASE.append((kode, nama, []))
    else:
        FASE[-1][2].append((kode, nama))


def diagram_wbs():
    jumlah_wp = sum(len(f[2]) for f in FASE)
    judul_tabel(f"Gambar 3.1 Diagram WBS (3 level, {len(FASE)} fase, {jumlah_wp} work package)")
    t = doc.add_table(rows=3, cols=len(FASE))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    puncak = t.cell(0, 0).merge(t.cell(0, len(FASE) - 1))
    r = puncak.paragraphs[0].add_run("1  Sistem Peringatan Dini Cuaca Publik (SPDC)")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    arsir(puncak, "1F4E79")
    for j, (kode, nama, wps) in enumerate(FASE):
        isi_runs(t.cell(1, j).paragraphs[0], f"**{kode}**\n**{nama}**", 10)
        arsir(t.cell(1, j), "BDD7EE")
        sel_wp = t.cell(2, j)
        for n, (kode_wp, nama_wp) in enumerate(wps):
            par = sel_wp.paragraphs[0] if n == 0 else sel_wp.add_paragraph()
            isi_runs(par, f"**{kode_wp}** {nama_wp}", 9.5)
            par.paragraph_format.space_after = Pt(6)
    atur_lebar(t, [int(lebar_isi() / len(FASE))] * len(FASE))
    for baris_t in t.rows[:2]:
        for c in baris_t.cells:
            for par in c.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.space_before = Pt(2)
    for c in t.rows[2].cells:
        for par in c.paragraphs:
            par.paragraph_format.line_spacing = 1.0
    p("Baris atas = level 1 (proyek), baris tengah = level 2 (fase), kolom bawah = level 3 (*work package*). "
      "WBS menunjukkan pemecahan pekerjaan, bukan urutan; urutan pekerjaan ditunjukkan oleh jadwal (Bab 4).",
      align=WD_ALIGN_PARAGRAPH.CENTER, ukuran=10)


def level(nilai):
    return "Kritikal" if nilai >= 20 else "Tinggi" if nilai >= 13 else "Sedang" if nilai >= 6 else "Rendah"


def gantt():
    T = max(a["ef"] for a in AKTIVITAS)
    W = math.ceil(T / 5)
    judul_tabel("Gambar 4.1 Gantt Chart per minggu (merah = jalur kritis, biru = memiliki float, ◆ = milestone)")
    t = doc.add_table(rows=1, cols=2 + W)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    kepala = ["Kode", "Aktivitas"] + [str(i) for i in range(1, W + 1)]
    for i, h in enumerate(kepala):
        c = t.rows[0].cells[i]
        isi_runs(c.paragraphs[0], f"**{h}**", 7.5)
        arsir(c, "E8E8E8")
    for a in AKTIVITAS:
        cells = t.add_row().cells
        isi_runs(cells[0].paragraphs[0], a["kode"], 7.5)
        isi_runs(cells[1].paragraphs[0], a["nama"], 7.5)
        for w in range(W):
            awal, akhir = w * 5, w * 5 + 5
            c = cells[2 + w]
            if a["dur"] == 0:
                if awal < a["ef"] <= akhir:
                    isi_runs(c.paragraphs[0], "**◆**", 8)
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif max(awal, a["es"]) < min(akhir, a["ef"]):
                arsir(c, WARNA_KRITIS if a["fl"] == 0 else WARNA_FLOAT)
    lebar = lebar_isi()
    k_kode, k_nama = Cm(1.3), Cm(7.2)
    k_minggu = int((lebar - k_kode - k_nama) / W)
    atur_lebar(t, [k_kode, k_nama] + [k_minggu] * W)
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
    ulang_header(t)
    p("Satu kolom mewakili satu minggu (5 hari kerja); aktivitas pendek yang melewati pergantian minggu tampil di dua "
      "kolom. Gantt versi harian tersedia di berkas ClickUp.", ukuran=10)


def matriks(residual):
    nomor, ket = ("5.2", "setelah respons (residual)") if residual else ("5.1", "sebelum respons")
    judul_tabel(f"Gambar {nomor} Probability-Impact Matrix {ket}")
    t = doc.add_table(rows=6, cols=6)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    isi_runs(t.cell(0, 0).paragraphs[0], "**K ↓ / D →**", 9)
    arsir(t.cell(0, 0), "E8E8E8")
    for d in range(1, 6):
        isi_runs(t.cell(0, d).paragraphs[0], f"**{d}**", 9); arsir(t.cell(0, d), "E8E8E8")
    for baris, k in enumerate(range(5, 0, -1), start=1):
        isi_runs(t.cell(baris, 0).paragraphs[0], f"**{k}**", 9); arsir(t.cell(baris, 0), "E8E8E8")
        for d in range(1, 6):
            c = t.cell(baris, d)
            nilai = k * d
            ids = [r["id"] for r in RISIKO if (r["k2"], r["d2"]) == (k, d)] if residual else \
                  [r["id"] for r in RISIKO if (r["k"], r["d"]) == (k, d)]
            isi_runs(c.paragraphs[0], f"{nilai} ({level(nilai)[0]})", 9)
            if ids:
                isi_runs(c.add_paragraph(), "**" + ", ".join(ids) + "**", 8.5)
            arsir(c, WARNA_LEVEL[level(nilai)])
    atur_lebar(t, [Cm(2.4)] + [Cm(3.0)] * 5)
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
    p("R = Rendah (1–5), S = Sedang (6–12), T = Tinggi (13–19), K = Kritikal (20–25).",
      align=WD_ALIGN_PARAGRAPH.CENTER, ukuran=10)


# ================= SAMPUL =================
m = re.search(r"<!-- SAMPUL\n(.*?)-->", teks_md, re.S)
SAMPUL = dict(b.split(": ", 1) for b in m.group(1).strip().splitlines())
for _ in range(2):
    doc.add_paragraph()
par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
par.add_run().add_picture(str((DRAF / SAMPUL["logo"]).resolve()), width=Cm(4.5))
doc.add_paragraph()
t = doc.add_paragraph(style="Title"); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.add_run(SAMPUL["judul"])
for kunci, ukuran in (("subjudul", 13), ("jenis", 13)):
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(SAMPUL[kunci]); r.font.size = Pt(ukuran)
for _ in range(3):
    doc.add_paragraph()
p("Disusun oleh:", align=WD_ALIGN_PARAGRAPH.CENTER)
anggota = [a.split("|") for a in SAMPUL["anggota"].split("; ")]
ta = doc.add_table(rows=len(anggota), cols=2)
ta.alignment = WD_TABLE_ALIGNMENT.CENTER
ta.autofit = False
for i, (nama, nim) in enumerate(anggota):
    isi_runs(ta.cell(i, 0).paragraphs[0], nama)
    isi_runs(ta.cell(i, 1).paragraphs[0], nim)
atur_lebar(ta, [Cm(6.5), Cm(3.5)])
for _ in range(3):
    doc.add_paragraph()
for kunci in ("dosen", "prodi", "kampus", "waktu"):
    par = p(SAMPUL[kunci], align=WD_ALIGN_PARAGRAPH.CENTER)
    par.paragraph_format.space_after = Pt(2)
    if kunci in ("prodi", "kampus"):
        for r in par.runs:
            r.bold = True

# Isi di section baru (sampul tanpa kepala halaman dan nomor halaman)
isi = doc.add_section(WD_SECTION.NEW_PAGE)
isi.header.is_linked_to_previous = False
isi.footer.is_linked_to_previous = False
kepala_halaman(isi, SAMPUL["header"])
nomor_halaman(isi, mulai=1)

# ================= ISI (dari .md) =================
teks_isi = re.sub(r"<!-- SAMPUL\n.*?-->", "", teks_md, flags=re.S)
baris = teks_isi.splitlines()
i = 0
info_tabel = None      # (judul, persen) dari penanda "tabel:"
paragraf = []
baru_ganti_halaman = True   # awal section isi = halaman baru
h1_pertama = True


def tutup_paragraf():
    global paragraf, baru_ganti_halaman
    if paragraf:
        p(" ".join(paragraf))
        paragraf = []
        baru_ganti_halaman = False


while i < len(baris):
    b = baris[i].rstrip()
    s = b.strip()
    if s.startswith("<!--"):
        tutup_paragraf()
        komentar = s
        while not komentar.endswith("-->"):
            i += 1
            komentar += "\n" + baris[i]
        isi_k = komentar[4:-3].strip()
        if isi_k in ("LANDSCAPE", "PORTRAIT"):
            section_baru(isi_k == "LANDSCAPE")
            baru_ganti_halaman = True
        elif isi_k == "GANTT":
            gantt(); baru_ganti_halaman = False
        elif isi_k == "DIAGRAM_WBS":
            diagram_wbs(); baru_ganti_halaman = False
        elif isi_k in ("MATRIKS_AWAL", "MATRIKS_RESIDUAL"):
            matriks(isi_k == "MATRIKS_RESIDUAL"); baru_ganti_halaman = False
        elif isi_k.startswith("tabel:"):
            bag = dict(x.strip().split(":", 1) for x in isi_k.split("|"))
            judul = bag.get("tabel", "").strip() or None
            persen = [float(x) for x in bag["lebar"].split(",")] if "lebar" in bag else None
            info_tabel = (judul, persen)
        i += 1
        continue
    if not s:
        tutup_paragraf(); i += 1; continue
    if s.startswith("#"):
        tutup_paragraf()
        tingkat = len(s) - len(s.lstrip("#"))
        teks = s[tingkat:].strip()
        if tingkat == 1:
            if not baru_ganti_halaman and not h1_pertama:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h1_pertama = False
        h = doc.add_heading("", level=tingkat)
        isi_runs(h, teks)
        h.paragraph_format.keep_with_next = True
        baru_ganti_halaman = False
        i += 1
        continue
    if s.startswith("|"):
        tutup_paragraf()
        blok = []
        while i < len(baris) and baris[i].startswith("|"):
            blok.append(sel(baris[i])); i += 1
        judul, persen = info_tabel or (None, None)
        tabel(blok[0], blok[2:], persen, judul)
        info_tabel = None
        baru_ganti_halaman = False
        continue
    mb = re.match(r"^- (.*)", s)
    mn = re.match(r"^(\d+)\. (.*)", s)
    if mb or mn:
        tutup_paragraf()
        butir(mb.group(1)) if mb else butir(mn.group(2), nomor=mn.group(1))
        baru_ganti_halaman = False
        i += 1
        continue
    paragraf.append(s)
    i += 1
tutup_paragraf()

HASIL.mkdir(exist_ok=True)
doc.save(OUT)
print("Tersimpan:", OUT)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HASIL), str(OUT)],
               check=True, capture_output=True)
print("Tersimpan:", OUT.with_suffix(".pdf"))
